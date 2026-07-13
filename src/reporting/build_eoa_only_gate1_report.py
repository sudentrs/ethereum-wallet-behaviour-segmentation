from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "reports" / "eoa_only_gate1_clustering_report.docx"


def pct(value: float) -> str:
    return f"{value * 100:.2f}%"


def num(value: float | int) -> str:
    return f"{int(round(float(value))):,}"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9DEE8") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        set_col_widths(table, widths)
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if row_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    style_table(table, widths)
    doc.add_paragraph()
    return table


def add_kv_table(doc, rows):
    return add_table(doc, ["Measure", "Value"], rows, widths=[3.2, 3.1])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("EOA-only Gate 1 report")
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(89, 89, 89)


def add_title(doc):
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("EOA-Only Ethereum Wallet Segmentation")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Gate 1 contract-contamination result, revised clustering, role interpretation, and k-sensitivity")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_callout(doc, title, text, fill="F4F6F9"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 77, 120)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        run.font.size = Pt(10)
    doc.add_paragraph()


def build_report():
    contamination = pd.read_csv(ROOT / "results/eoa_contract_contamination_full/cluster_contract_contamination_summary.csv")
    eoa_summary = pd.read_csv(ROOT / "data/processed/wallet_kmeans6_expanded_daily_features_eoa_only_summary.csv").iloc[0]
    kmeans_summary = pd.read_csv(ROOT / "results/eoa_only_kmeans6/kmeans_summary.csv").iloc[0]
    sizes = pd.read_csv(ROOT / "results/eoa_only_kmeans6/cluster_sizes.csv")
    roles = pd.read_csv(ROOT / "results/eoa_only_kmeans6/cluster_descriptions/eoa_only_cluster_behavior_labels.csv")
    correspondence = pd.read_csv(ROOT / "results/eoa_only_kmeans6/mixed_comparison/mixed_to_eoa_only_comparison_summary.csv").iloc[0]
    matches = pd.read_csv(ROOT / "results/eoa_only_kmeans6/mixed_comparison/mixed_to_eoa_only_hungarian_matches.csv")
    k_summary = pd.read_csv(ROOT / "results/eoa_only_kmeans6/k_sensitivity_full/k_sensitivity_summary.csv")

    doc = Document()
    configure_document(doc)
    add_title(doc)

    doc.add_heading("Executive Summary", level=1)
    add_callout(
        doc,
        "Main conclusion",
        "The original mixed-address KMeans6 result should now be treated as a pre-Gate-1 baseline. "
        "Full EOA/contract classification showed that mixed cluster C3 is 47.95% contracts, above the 15% Gate 1 threshold. "
        "After removing contracts, 2.44 million EOA wallets remain and support a revised wallet-behaviour clustering branch.",
    )
    doc.add_paragraph(
        "The EOA-only KMeans6 result gives an interpretable six-role vocabulary, but the full k-sensitivity analysis shows that "
        "k=6 is not the most stable EOA-only resolution. Coarser k=2 and k=3 solutions are strongest by silhouette and seed stability, "
        "while k=5 is more stable than k=6 among richer multi-role segmentations."
    )

    doc.add_heading("1. What Was Done", level=1)
    for text in [
        "Classified all original KMeans6 wallet addresses as EOA or contract using Ethereum bytecode checks.",
        "Measured contract contamination per original mixed-address cluster.",
        "Filtered the expanded wallet feature table to EOA-only wallets.",
        "Re-scaled the EOA-only feature table using the same robust transformation logic.",
        "Ran a new EOA-only KMeans6 clustering and labelled the resulting behavioural roles.",
        "Compared old mixed-address roles against new EOA-only roles using overlap and Hungarian matching.",
        "Ran full-data k-sensitivity and seed-stability checks for k=2 through k=10.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("2. Gate 1 Contract-Contamination Result", level=1)
    doc.add_paragraph(
        "Gate 1 asked whether any behavioural cluster was materially contaminated by smart contracts. The working threshold was 15%. "
        "The full classification confirms that mixed C3 fails this gate decisively."
    )
    rows = []
    for _, row in contamination.iterrows():
        rows.append([
            f"C{int(row['cluster'])}",
            num(row["sampled_addresses"]),
            num(row["eoa_count"]),
            num(row["contract_count"]),
            pct(row["contract_share"]),
        ])
    add_table(doc, ["Mixed cluster", "Total", "EOA", "Contract", "Contract share"], rows, widths=[1.05, 1.35, 1.35, 1.35, 1.25])
    doc.add_paragraph(
        "Interpretation: mixed C3 is nearly half contract addresses. C1 is also close to the threshold at 14.32%, so it should be discussed carefully even though it does not formally fail Gate 1."
    )

    doc.add_heading("3. EOA-Only Dataset", level=1)
    add_kv_table(
        doc,
        [
            ["Original mixed-address wallets", num(eoa_summary["feature_wallets"])],
            ["EOA-only wallets retained", num(eoa_summary["eoa_wallets"])],
            ["Contract addresses removed", num(eoa_summary["contract_wallets"])],
            ["EOA share", pct(eoa_summary["eoa_share_of_classified"])],
            ["Contract share", pct(eoa_summary["contract_share_of_classified"])],
        ],
    )
    doc.add_paragraph(
        "This EOA-only branch is the appropriate branch for paper-facing claims about wallet behaviour. The old mixed branch remains useful as a baseline and as evidence that contract contamination changes the role vocabulary."
    )

    doc.add_heading("4. EOA-Only KMeans6 Result", level=1)
    add_kv_table(
        doc,
        [
            ["Wallets clustered", num(kmeans_summary["wallets_clustered"])],
            ["Feature count", num(kmeans_summary["feature_count"])],
            ["k", num(kmeans_summary["k"])],
            ["Random seed", num(kmeans_summary["random_state"])],
            ["Silhouette", f"{kmeans_summary['silhouette']:.4f}"],
            ["Inertia", num(kmeans_summary["inertia"])],
        ],
    )

    size_rows = [[f"C{int(r.cluster)}", num(r.wallet_count), pct(r.wallet_share)] for r in sizes.itertuples()]
    add_table(doc, ["EOA-only cluster", "Wallets", "Share"], size_rows, widths=[2.0, 2.0, 2.0])

    doc.add_heading("5. Revised EOA-Only Behavioural Roles", level=1)
    doc.add_paragraph(
        "The following role labels summarize median robust-scaled feature profiles after contract addresses were removed. They should not be read as verified real-world identities."
    )
    role_rows = []
    for _, r in roles.iterrows():
        role_rows.append([
            f"C{int(r['cluster'])}",
            pct(r["wallet_share"]),
            r["behavior_label"],
            {
                0: "Low value-flow wallets with weak total inflow/outflow and low flow volatility.",
                1: "Directionally imbalanced wallets with positive net-flow features but relatively low transaction intensity.",
                2: "High-activity wallets with persistent activity, large value flow, and high volatility.",
                3: "Wallets with concentrated transaction/value movement, especially outflow-related activity, but lower persistence.",
                4: "Short-lived but very dense activity wallets with strong transaction concentration and outflow leaning.",
                5: "Wallets with elevated net-flow, total-flow volatility, and persistent but irregular value movement.",
            }[int(r["cluster"])],
        ])
    add_table(doc, ["Cluster", "Share", "Role label", "Interpretation"], role_rows, widths=[0.75, 0.85, 2.1, 2.65])

    doc.add_heading("6. How The New Roles Differ From The Old Mixed Roles", level=1)
    doc.add_paragraph(
        "The new EOA-only clustering is related to the mixed-address clustering, but it is not simply a relabelling. Contract removal reorganizes the role structure."
    )
    add_kv_table(
        doc,
        [
            ["Overlapping EOA wallets", num(correspondence["overlapping_eoa_wallets"])],
            ["Hungarian matched wallets", num(correspondence["matched_wallets_hungarian"])],
            ["Hungarian matched share", pct(correspondence["matched_share_hungarian"])],
            ["Adjusted Rand Index", f"{correspondence['adjusted_rand_index']:.4f}"],
            ["Normalized Mutual Information", f"{correspondence['normalized_mutual_information']:.4f}"],
        ],
    )
    match_rows = []
    for _, r in matches.iterrows():
        match_rows.append([
            f"Mixed C{int(r['mixed_cluster'])}",
            f"EOA C{int(r['eoa_only_cluster'])}",
            num(r["matched_wallets"]),
            pct(r["share_of_mixed_cluster_eoas"]),
        ])
    add_table(doc, ["Old role", "New role", "Matched EOA wallets", "Share of old-role EOAs"], match_rows, widths=[1.25, 1.25, 1.8, 1.95])
    doc.add_paragraph(
        "The most important transition is mixed C3. Although mixed C3 was heavily contract-contaminated, its EOA subset is not random: 84.52% of mixed C3 EOAs map to EOA-only C4. This means the old C3 contained a real EOA behavioural signal, but the pre-Gate-1 interpretation was blended with a large contract population."
    )

    doc.add_heading("7. K Sensitivity And Stability", level=1)
    doc.add_paragraph(
        "The full EOA-only k-sensitivity run tested k=2 through k=10 on all 2,443,868 EOA wallets, using seeds 492, 493, and 494. Silhouette was estimated on 50,000 wallets for tractability."
    )
    compact_k = k_summary[["k", "silhouette_mean", "stability_ari_mean", "stability_nmi_mean", "min_cluster_share_mean", "max_cluster_share_mean"]]
    k_rows = []
    for _, r in compact_k.iterrows():
        k_rows.append([
            int(r["k"]),
            f"{r['silhouette_mean']:.3f}",
            f"{r['stability_ari_mean']:.3f}",
            f"{r['stability_nmi_mean']:.3f}",
            pct(r["min_cluster_share_mean"]),
            pct(r["max_cluster_share_mean"]),
        ])
    add_table(doc, ["k", "Silhouette", "Stability ARI", "Stability NMI", "Min share", "Max share"], k_rows, widths=[0.55, 1.05, 1.25, 1.25, 1.1, 1.1])
    add_callout(
        doc,
        "K-selection implication",
        "k=2 and k=3 are strongest by metric-only criteria but are coarse. Among richer multi-role segmentations, k=5 is more stable than k=6. Therefore, k=6 can be presented as an interpretable fine-grained vocabulary, but not as the most stable EOA-only solution without further downstream validation.",
        fill="FFF8E6",
    )

    doc.add_heading("8. Recommended Interpretation For The Paper", level=1)
    for text in [
        "Use the mixed-address KMeans6 result only as a pre-Gate-1 baseline.",
        "State that Gate 1 failed because mixed C3 is 47.95% contracts.",
        "Use the EOA-only branch for wallet-behaviour claims.",
        "Avoid saying k=6 is the true number of wallet roles.",
        "Describe k=6 as an interpretable fine-grained EOA-only role vocabulary.",
        "Mention that k=5 is more stable among multi-role EOA-only segmentations, while k=2 and k=3 capture coarser high-level structure.",
        "Run downstream validation on the EOA-only branch before final submission: transaction-graph validation, temporal transition analysis, and stability/subsampling checks.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Appendix: Key Output Files", level=1)
    add_table(
        doc,
        ["File", "Purpose"],
        [
            ["data/processed/wallet_kmeans6_expanded_daily_features_eoa_only.csv", "EOA-only expanded feature table."],
            ["data/processed/wallet_kmeans6_expanded_daily_features_eoa_only_robust_scaled.csv", "EOA-only robust-scaled feature table."],
            ["results/eoa_only_kmeans6/wallet_eoa_only_kmeans_clusters.csv", "EOA-only KMeans6 assignments."],
            ["results/eoa_only_kmeans6/cluster_descriptions/eoa_only_cluster_behavior_labels.csv", "Role labels and profile summaries."],
            ["results/eoa_only_kmeans6/mixed_comparison/", "Old-vs-new cluster correspondence outputs."],
            ["results/eoa_only_kmeans6/k_sensitivity_full/k_sensitivity_summary.csv", "Full-data k-sensitivity and seed-stability table."],
        ],
        widths=[3.6, 2.75],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
