from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_final_report
from build_cluster_stability_complete_report import (
    ROOT,
    add_bullets,
    add_dataframe_table,
    add_para,
    load_csv,
)


OUT = ROOT / "ENS491-492_Final_Report_Full_Length_Template_Updated.docx"


def find_heading(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"Heading not found: {text}")


def insert_doc_before(doc, heading_text, fragment):
    target = find_heading(doc, heading_text)
    target_el = target._p
    body = fragment._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        target_el.addprevious(deepcopy(child))


def insert_doc_after_heading(doc, heading_text, fragment):
    target = find_heading(doc, heading_text)
    anchor = target._p
    body = fragment._body._element
    for child in reversed(list(body)):
        if child.tag.endswith("sectPr"):
            continue
        anchor.addnext(deepcopy(child))


def insert_figure_before(doc, heading_text, rel_path, caption, width=6.0):
    target = find_heading(doc, heading_text)
    path = ROOT / rel_path
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(9)
    target._p.addprevious(p._p)
    target._p.addprevious(cap._p)


def methodology_fragment():
    frag = Document()
    distance_consensus = load_csv("paper_process_outputs/distance_metric_consensus_matches.csv")
    corr_perm = load_csv("paper_process_outputs/cross_dataset_correspondence_permutation_summary.csv")
    feature_corr = load_csv("paper_process_outputs/cross_feature_set_spearman_summary.csv")
    obs = load_csv("paper_process_outputs/minimum_observation_experiment_summary.csv")
    temporal = load_csv("paper_process_outputs/paper_ready_temporal_stability_summary.csv")
    state_retention = load_csv("paper_process_outputs/paper_ready_state_retention_rates.csv")

    frag.add_heading("Advisor-Guided Paper-Process Validation", level=2)
    add_para(
        frag,
        "After the first full final-report draft, the advisor report clarified which additional analyses were needed to make the cluster-stability argument paper-ready. These additions do not replace the original KMeans6, network-science, transaction-network, temporal, and internal-subcluster work. Instead, they formalize and pressure-test the claims that were already emerging from those analyses.",
    )
    add_para(
        frag,
        "The first addition was a formal cluster correspondence rule. Cluster labels from two different analyses cannot be compared by their numeric names alone. For example, Cluster 1 in one analysis does not automatically represent the same behaviour as Cluster 1 in another analysis. Therefore, the correspondence problem was defined as a maximum bipartite matching problem: clusters from one analysis and clusters from the other analysis are treated as two sets, and the algorithm finds the best one-to-one pairing under a chosen distance metric.",
    )
    add_para(
        frag,
        "The project then compared cluster profiles using four distance metrics: cosine, weighted Euclidean, Mahalanobis, and Wasserstein. Each metric captures similarity differently. Cosine distance focuses on the direction or pattern of the feature profile. Weighted Euclidean distance measures overall numeric closeness while allowing feature families to have different importance. Mahalanobis distance accounts for covariance among features. Wasserstein distance provides a distribution-oriented comparison. A cluster pair was treated as a core correspondence when it appeared under at least three of the four metrics.",
    )
    add_dataframe_table(frag, distance_consensus, max_rows=10, font_size=8)
    add_para(
        frag,
        "The second addition was a cross-dataset correspondence permutation test. This test asks whether the observed wallet-label alignment between two clusterings is stronger than what would be expected under random labels. The observed maximum bipartite overlap was compared against 1,000 randomized label permutations. This is separate from the earlier feature-graph random-label test: the graph test validates label alignment with feature-similarity structure, while this new test validates cross-dataset label alignment.",
    )
    add_dataframe_table(
        frag,
        corr_perm[
            [
                "n_wallets",
                "observed_max_bipartite_overlap",
                "null_mean",
                "null_95pct_lower",
                "null_95pct_upper",
                "z_score",
                "empirical_right_tail_p_value",
            ]
        ],
        max_rows=3,
        font_size=8,
    )
    add_para(
        frag,
        "The third addition was a cross-feature-set Spearman correlation analysis. The purpose was to test whether the original sparse feature set and the expanded daily feature set were genuinely different. If they were almost identical, then convergence across analyses would be less meaningful. The Spearman analysis showed that the two feature sets are related but not redundant.",
    )
    add_dataframe_table(frag, feature_corr, max_rows=3, font_size=8)
    add_para(
        frag,
        "The fourth addition was a minimum observation experiment. The sparse interval dataset was subsampled at 10%, 25%, 50%, 75%, and 100%. For each retained fraction, clustering was rerun and compared with the existing reference labels. This tests whether broad behavioural roles can be recovered when less data is available and whether finer roles require denser observation.",
    )
    add_dataframe_table(frag, obs, max_rows=10, font_size=8)
    add_para(
        frag,
        "Finally, the existing temporal results were repackaged into paper-ready stability summaries. The original temporal analysis already assigned wallets to monthly states and computed transition probabilities. The new summary files report the number of consecutive wallet-month pairs, the mean self-transition rate, and state-level retention rates in a form that can be inserted directly into a paper results section.",
    )
    add_dataframe_table(frag, temporal, max_rows=3, font_size=8)
    add_dataframe_table(frag, state_retention, max_rows=10, font_size=8)
    return frag


def discussion_fragment():
    frag = Document()
    corr_perm = load_csv("paper_process_outputs/cross_dataset_correspondence_permutation_summary.csv")
    feature_corr = load_csv("paper_process_outputs/cross_feature_set_spearman_summary.csv")
    obs = load_csv("paper_process_outputs/minimum_observation_experiment_summary.csv")

    frag.add_heading("Advisor-Guided Stability Findings", level=2)
    add_para(
        frag,
        f"The advisor-guided analyses strengthen the original discussion by showing that the cluster-stability claim can be tested more formally. The distance-metric consensus results identify five core profile correspondences across the compared analyses. The separate wallet-label correspondence test produced an observed maximum bipartite overlap of {corr_perm.iloc[0]['observed_max_bipartite_overlap']:.3f}, compared with a null mean of {corr_perm.iloc[0]['null_mean']:.3f}, with p = {corr_perm.iloc[0]['empirical_right_tail_p_value']:.6f}. This indicates statistically significant cross-dataset alignment.",
    )
    add_para(
        frag,
        "This result should be interpreted carefully. Profile-level correspondence and wallet-level label overlap answer different questions. Profile correspondence asks whether behavioural archetypes reappear across analyses. Wallet-level overlap asks whether the same addresses receive aligned labels across two clusterings. The first supports the stability of archetypes; the second shows that direct address-level alignment is significantly above chance but more moderate in magnitude.",
    )
    add_para(
        frag,
        f"The cross-feature-set Spearman analysis also strengthens the robustness argument. Across {int(feature_corr.iloc[0]['sampled_wallets']):,} sampled overlapping wallets, the median absolute correlation was {feature_corr.iloc[0]['median_abs_spearman']:.3f}, and only {100*feature_corr.iloc[0]['share_abs_corr_ge_0_8']:.1f}% of feature pairs had absolute correlation at or above 0.8. This means the feature spaces share behavioural information but are not simply copies of each other.",
    )
    add_para(
        frag,
        "The minimum observation experiment gives a practical interpretation of stability. It suggests that some broad roles can be recovered from reduced data, while finer-grained roles are more sensitive to observation density. This finding is useful for future research because real blockchain datasets often differ in temporal coverage, sampling density, and feature availability.",
    )
    return frag


def conclusion_fragment():
    frag = Document()
    frag.add_heading("Updated Paper-Process Conclusion", level=2)
    add_para(
        frag,
        "The advisor-guided additions change the status of the work from a strong project report into a more paper-ready empirical validation study. The original final report showed that the wallet roles are interpretable, statistically separable, graph-theoretically meaningful, transaction-network relevant, temporally dynamic, and internally heterogeneous. The new additions show that the correspondence claim can be formalized, tested against a null model, checked against feature redundancy, and evaluated under reduced observation density.",
    )
    add_para(
        frag,
        "The most defensible final claim is therefore not that every wallet always keeps the same exact label across every analysis. The stronger claim is that Ethereum wallet populations contain recurring behavioural structures that can be discovered, interpreted, and validated across several analytical views. These structures provide a stronger empirical foundation for the original AI-governed crypto-finance simulation idea.",
    )
    return frag


def appendix_fragment():
    frag = Document()
    frag.add_heading("Appendix C: Advisor-Guided Paper-Process Outputs", level=2)
    rows = [
        ["paper_process_outputs/distance_metric_consensus_matches.csv", "Core cluster correspondences across four distance metrics."],
        ["paper_process_outputs/cross_dataset_correspondence_permutation_summary.csv", "Permutation test for cross-dataset wallet-label correspondence."],
        ["paper_process_outputs/cross_feature_set_spearman_summary.csv", "Summary of Spearman correlation between original and expanded feature sets."],
        ["paper_process_outputs/minimum_observation_experiment_summary.csv", "Observation-density subsampling experiment."],
        ["paper_process_outputs/paper_ready_temporal_stability_summary.csv", "Paper-ready summary of individual-level monthly stability."],
        ["run_paper_process_remaining_analyses.py", "Reproducible script that generated the new paper-process outputs."],
    ]
    import pandas as pd

    add_dataframe_table(frag, pd.DataFrame(rows, columns=["Artifact", "Purpose"]), max_rows=20, font_size=8)
    return frag


def main():
    build_final_report.OUT = OUT
    build_final_report.build()
    doc = Document(OUT)
    insert_doc_before(doc, "DISCUSSION", methodology_fragment())
    insert_figure_before(
        doc,
        "DISCUSSION",
        "paper_process_outputs/cross_dataset_correspondence_permutation_null.png",
        "Figure M1. Cross-dataset cluster correspondence permutation test.",
        width=5.9,
    )
    insert_figure_before(
        doc,
        "DISCUSSION",
        "paper_process_outputs/cross_feature_set_spearman_heatmap.png",
        "Figure M2. Spearman correlation heatmap between the original and expanded feature sets.",
        width=6.3,
    )
    insert_figure_before(
        doc,
        "DISCUSSION",
        "paper_process_outputs/minimum_observation_correspondence_curve.png",
        "Figure M3. Minimum observation experiment showing correspondence under sparse-data subsampling.",
        width=5.9,
    )
    insert_figure_before(
        doc,
        "DISCUSSION",
        "paper_process_outputs/paper_ready_monthly_retention_rates.png",
        "Figure M4. Month-to-month self-transition rates for wallet behaviour states.",
        width=6.2,
    )
    insert_doc_before(doc, "Scientific, Technological, and Socio-Economic Impacts", discussion_fragment())
    insert_doc_before(doc, "APPENDIX", conclusion_fragment())
    insert_doc_before(doc, "REFERENCES", appendix_fragment())
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
