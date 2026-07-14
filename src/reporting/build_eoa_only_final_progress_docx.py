from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE_MD = ROOT / "docs" / "EOA_ONLY_FINAL_PROGRESS_REPORT.md"
OUT_DOCX = ROOT / "reports" / "EOA_Only_Final_Progress_Report.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
LIGHT_GRAY = "F2F4F7"
BORDER_GRAY = "D9E2EC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER_GRAY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def style_table(table, widths=None) -> None:
    if widths:
        set_table_width(table, widths)
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
            if r == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def set_doc_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("EOA-Only Ethereum Wallet Behaviour Segmentation")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(11, 37, 69)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Final progress report from contract removal through K5 academic strengthening")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(85, 85, 85)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("Prepared from the project repository analysis outputs.")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(85, 85, 85)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    set_cell_border(cell, "C8D3DF")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(31, 58, 95)
    run.font.size = Pt(10.5)
    doc.add_paragraph()


def parse_inline(text: str, paragraph) -> None:
    parts = text.split("`")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.name = "Consolas" if idx % 2 else "Calibri"
        run.font.size = Pt(9.5 if idx % 2 else 11)
        if idx % 2:
            run.font.color.rgb = RGBColor(31, 77, 120)


def parse_markdown_table(lines):
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|") or not stripped.endswith("|"):
            return None
        cells = [c.strip() for c in stripped.strip("|").split("|")]
        if all(set(c.replace(":", "").strip()) <= {"-"} for c in cells):
            continue
        rows.append(cells)
    return rows if rows else None


def add_markdown_table(doc: Document, rows) -> None:
    if not rows:
        return
    max_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(max_cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = row[c_idx] if c_idx < len(row) else ""
    if max_cols <= 3:
        widths = [1.5, 2.4, 2.4][:max_cols]
    elif max_cols == 4:
        widths = [0.8, 2.0, 1.4, 2.3]
    elif max_cols == 5:
        widths = [0.6, 1.2, 1.2, 1.4, 2.1]
    else:
        widths = [6.5 / max_cols] * max_cols
    style_table(table, widths)
    doc.add_paragraph()


def add_csv_table(doc: Document, title: str, csv_path: Path, columns=None, widths=None) -> None:
    if not csv_path.exists():
        return
    doc.add_heading(title, level=3)
    df = pd.read_csv(csv_path)
    if columns:
        df = df[columns]
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_float_dtype(df[col]):
            df[col] = df[col].map(lambda x: f"{x:.4f}" if pd.notna(x) else "")
    table = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
    table.style = "Table Grid"
    for c_idx, col in enumerate(df.columns):
        table.cell(0, c_idx).text = str(col).replace("_", " ")
    for r_idx, (_, row) in enumerate(df.iterrows(), start=1):
        for c_idx, col in enumerate(df.columns):
            table.cell(r_idx, c_idx).text = str(row[col])
    style_table(table, widths)
    doc.add_paragraph()


def build_docx() -> None:
    doc = Document()
    set_doc_styles(doc)
    add_title(doc)

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            # Title already handled.
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("> "):
            quote_lines = []
            while i < len(lines) and lines[i].startswith("> "):
                quote_lines.append(lines[i][2:].strip())
                i += 1
            add_callout(doc, " ".join(quote_lines))
            continue
        if line.startswith("- "):
            while i < len(lines) and lines[i].startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                parse_inline(lines[i][2:].strip(), p)
                i += 1
            continue
        if line[0:2].isdigit() and ". " in line[:5]:
            p = doc.add_paragraph(style="List Number")
            parse_inline(line.split(". ", 1)[1], p)
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_markdown_table(table_lines)
            add_markdown_table(doc, rows)
            continue
        p = doc.add_paragraph()
        parse_inline(line, p)
        i += 1

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("Appendix: Key Generated Tables", level=1)
    add_csv_table(
        doc,
        "K4/K5/K6 Downstream Comparison",
        ROOT / "results" / "eoa_k4_k5_k6_comparison" / "eoa_k4_k5_k6_downstream_summary.csv",
        columns=[
            "k",
            "seed_stability_ari",
            "network_random_null_p",
            "temporal_self_transition_mean",
            "two_step_return_mean",
            "verdict",
        ],
        widths=[0.45, 1.0, 1.0, 1.1, 1.0, 1.95],
    )
    add_csv_table(
        doc,
        "K5 Per-Cluster Stability",
        ROOT / "results" / "k5_academic_strengthening" / "per_cluster_stability_table.csv",
        columns=[
            "cluster",
            "behavior_label",
            "wallet_share",
            "matched_recovery_mean",
            "self_transition_probability",
            "role_reliability_summary",
            "interpretation_caution",
        ],
        widths=[0.45, 1.55, 0.75, 0.95, 0.95, 0.85, 1.0],
    )
    add_csv_table(
        doc,
        "Trajectory Type Summary",
        ROOT / "results" / "k5_academic_strengthening" / "trajectory_type_summary.csv",
        columns=[
            "trajectory_type",
            "wallet_count",
            "mean_months_observed",
            "mean_switch_count",
            "mean_two_step_return_count",
            "wallet_share",
        ],
        widths=[1.45, 0.9, 1.1, 0.9, 1.25, 0.9],
    )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_docx()
