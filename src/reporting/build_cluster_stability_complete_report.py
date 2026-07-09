from pathlib import Path
import math
import textwrap

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Complete_Project_Findings_and_Paper_Process_Report.docx"


COLORS = {
    "ink": RGBColor(31, 41, 55),
    "muted": RGBColor(88, 99, 113),
    "accent": RGBColor(35, 91, 118),
    "rust": RGBColor(158, 83, 65),
    "green": RGBColor(62, 120, 104),
    "light": "EAF2F7",
    "line": "D7DEE8",
    "header": "DDEBF2",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header=True):
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9)
            if i == 0 and header:
                set_cell_shading(cell, COLORS["header"])
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = COLORS["ink"]


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_doc_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = COLORS["ink"]
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, before, after, color in [
        ("Title", 23, 0, 12, COLORS["accent"]),
        ("Subtitle", 12, 0, 10, COLORS["muted"]),
        ("Heading 1", 16, 14, 6, COLORS["accent"]),
        ("Heading 2", 13, 10, 4, COLORS["ink"]),
        ("Heading 3", 11, 8, 3, COLORS["rust"]),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.text = "Ethereum Wallet Behaviour Segmentation | Complete Findings and Paper Process"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.color.rgb = COLORS["muted"]
        footer = section.footer
        p = footer.paragraphs[0]
        add_page_number(p)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.color.rgb = COLORS["muted"]


def add_para(doc, text, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.name = "Arial"
        r.font.color.rgb = COLORS["ink"]
        remaining = text[len(bold_lead):] if text.startswith(bold_lead) else text
        if remaining:
            p.add_run(remaining)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Arial"
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = COLORS["ink"]


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.color.rgb = COLORS["muted"]


def add_figure(doc, path, caption, width=6.2):
    path = ROOT / path
    if not path.exists():
        add_para(doc, f"[Figure missing: {path.name}]", style=None)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_dataframe_table(doc, df, cols=None, max_rows=12, font_size=8.5):
    if cols:
        df = df[cols]
    df = df.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(df.columns))
    for j, col in enumerate(df.columns):
        set_cell_text(table.rows[0].cells[j], col, bold=True, size=font_size)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(df.columns):
            val = row[col]
            if isinstance(val, float):
                if abs(val) >= 1000:
                    text = f"{val:,.0f}"
                else:
                    text = f"{val:.4f}".rstrip("0").rstrip(".")
            else:
                text = val
            set_cell_text(cells[j], text, size=font_size)
    style_table(table)
    doc.add_paragraph()
    return table


def load_csv(rel):
    return pd.read_csv(ROOT / rel)


def pct(x):
    return f"{100*x:.1f}%"


def build_report():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    set_doc_styles(doc)
    add_header_footer(doc)

    # Data
    cluster_sizes = load_csv("dataverse_files/network_science_kmeans6_full/cluster_sizes.csv")
    behavior = load_csv("dataverse_files/network_science_kmeans6_full/cluster_behavior_interpretation.csv")
    feature_graph = load_csv("dataverse_files/network_science_kmeans6_full/feature_similarity_network_summary.csv")
    graph_perm = load_csv("dataverse_files/network_science_kmeans6_full/feature_graph_random_label_hypothesis_test.csv")
    alt_k = load_csv("dataverse_files/network_science_kmeans6_full/alternative_k_clusterability.csv")
    distance_consensus = load_csv("paper_process_outputs/distance_metric_consensus_matches.csv")
    corr_perm = load_csv("paper_process_outputs/cross_dataset_correspondence_permutation_summary.csv")
    temporal = load_csv("paper_process_outputs/paper_ready_temporal_stability_summary.csv")
    state_retention = load_csv("paper_process_outputs/paper_ready_state_retention_rates.csv")
    feature_corr = load_csv("paper_process_outputs/cross_feature_set_spearman_summary.csv")
    obs = load_csv("paper_process_outputs/minimum_observation_experiment_summary.csv")

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    r = p.add_run("Complete Project Findings and Paper-Process Report")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(24)
    r.font.color.rgb = COLORS["accent"]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Ethereum Wallet Behaviour Segmentation, Cluster Stability, and Cross-Scale Validation")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.color.rgb = COLORS["muted"]

    meta = doc.add_table(rows=4, cols=2)
    meta_data = [
        ("Project", "ENS 491-492 Graduation Project"),
        ("Core dataset", "Ethereum ERC-20 wallet transaction behaviour"),
        ("Main empirical unit", "Wallet-level behaviour profile and monthly behaviour state"),
        ("Report purpose", "Consolidate completed work, advisor-required paper analyses, findings, and remaining manuscript tasks"),
    ]
    for i, (k, v) in enumerate(meta_data):
        set_cell_text(meta.rows[i].cells[0], k, bold=True, color=COLORS["accent"], size=9.5)
        set_cell_text(meta.rows[i].cells[1], v, size=9.5)
    style_table(meta, header=False)

    add_para(
        doc,
        "This report follows the chronological development of the project and then mirrors the advisor report structure: opportunity and gap, missing analytical pieces, completed additions, results, discussion, limitations, and next manuscript actions.",
    )
    doc.add_page_break()

    # Executive summary
    doc.add_heading("Executive Summary", level=1)
    add_para(
        doc,
        "The project began as an AI-governed crypto-bank simulator, but the technical center shifted toward an empirical question: can real Ethereum wallets be organized into stable, interpretable behavioural roles that survive validation across feature space, network structure, and time?",
    )
    add_para(
        doc,
        "The completed work now supports that direction. A full-year KMeans6 segmentation over 2.74 million wallets produced six interpretable wallet roles. Graph validation showed strong non-random structure in feature-similarity space, with high intra-cluster edge concentration, modularity, and assortativity. Temporal analysis showed that monthly behaviour is moderately persistent rather than random. New paper-process analyses added the missing formal correspondence definition, a cross-dataset correspondence permutation test, cross-feature-set correlation evidence, and a minimum observation experiment.",
    )

    add_bullets(
        doc,
        [
            "Full KMeans6 population: 2,738,243 wallets across six wallet behaviour roles.",
            f"Feature-similarity graph validation: {int(feature_graph.iloc[0]['nodes']):,} nodes, {int(feature_graph.iloc[0]['edges']):,} edges, modularity {feature_graph.iloc[0]['cluster_label_modularity']:.3f}, assortativity {feature_graph.iloc[0]['cluster_assortativity']:.3f}.",
            f"Random-label graph test: observed same-cluster edge ratio {graph_perm.iloc[0]['observed_intra_ratio']:.3f}, null mean {graph_perm.iloc[0]['permutation_null_mean']:.3f}, p-value {graph_perm.iloc[0]['permutation_p_value_right_tail']:.6f}.",
            f"Cross-dataset wallet-label correspondence: observed optimal overlap {corr_perm.iloc[0]['observed_max_bipartite_overlap']:.3f}, null mean {corr_perm.iloc[0]['null_mean']:.3f}, p-value {corr_perm.iloc[0]['empirical_right_tail_p_value']:.6f}.",
            f"Temporal stability: {int(temporal.iloc[0]['total_consecutive_wallet_month_pairs']):,} consecutive wallet-month pairs; mean monthly self-transition {temporal.iloc[0]['mean_monthly_self_transition_rate']:.3f}.",
            f"Feature-set independence: median absolute Spearman correlation {feature_corr.iloc[0]['median_abs_spearman']:.3f} across sampled overlapping wallets, suggesting related but non-redundant feature spaces.",
        ],
    )

    doc.add_heading("Chronological Project Development", level=1)
    add_para(
        doc,
        "The project did not start as a paper about cluster stability. It moved there through a sequence of increasingly empirical decisions. The chronology matters because it explains why the current paper-process work is not disconnected from the original system idea.",
    )
    chronology = pd.DataFrame(
        [
            ["Initial concept", "AI-governed crypto-bank simulator", "Simulated depositors, borrowers, governance participants, credit scoring, LLM advisory, and explainability."],
            ["Motivation refined", "Need for realistic agent types", "A simulator cannot be convincing if user personas are invented without evidence from real blockchain behaviour."],
            ["Empirical pivot", "Wallet behaviour segmentation", "The project shifted toward discovering wallet roles from real Ethereum transaction data."],
            ["Core modelling", "KMeans6 full-year segmentation", "Wallet-level features were built and clustered into six interpretable behavioural roles."],
            ["Network validation", "Feature-similarity and transaction-network evidence", "Clusters were tested against graph structure, modularity, assortativity, and random-label null models."],
            ["Dynamic analysis", "Monthly behaviour-state transitions", "Wallets were assigned to monthly behavioural states to measure persistence and transitions."],
            ["Advisor paper process", "Formalize and pressure-test the claims", "New analyses added correspondence significance testing, feature-set correlation, and observation-density checks."],
        ],
        columns=["Stage", "Focus", "What changed or was completed"],
    )
    add_dataframe_table(doc, chronology, max_rows=10)

    doc.add_heading("1. Opportunity and Core Thesis", level=1)
    add_para(
        doc,
        "The advisor report frames the opportunity as an empirical cluster-stability paper. The central thesis is that behavioural stratification in on-chain wallet populations may be a structural property of the user distribution rather than an artifact of one dataset, one time window, one feature set, or one clustering run.",
    )
    add_para(
        doc,
        "The strongest contribution is not a new clustering algorithm. It is the layered validation of existing clustering output: cross-scale cluster correspondence, graph-based validation, temporal transition evidence, feature-set non-redundancy, and observation-density sensitivity. This matches the advisor's recommendation to treat the project as an empirical validation study with a formalized method.",
    )

    doc.add_heading("Research Questions", level=2)
    add_bullets(
        doc,
        [
            "Do behavioural archetypes discovered in a limited or sparse on-chain observation window reproduce in a larger longitudinal dataset?",
            "Are the discovered roles robust to different feature-engineering philosophies?",
            "Do the roles correspond to real graph structure rather than arbitrary mathematical partitions?",
            "How stable are wallet behaviours at the individual level from month to month?",
            "How much observation density is needed to recover the same core roles?",
        ],
    )

    doc.add_heading("2. Data, Features, and Cluster Construction", level=1)
    add_para(
        doc,
        "The project uses wallet-level Ethereum activity derived from ERC-20 transaction data. The full KMeans6 analysis uses expanded daily behavioural features such as active days, activity density, total transactions, total inflow and outflow, net flow, daily averages, maximum daily values, volatility measures, active-day directionality, and concentration ratios.",
    )
    add_para(
        doc,
        "A second comparison analysis uses a sparser transaction-structure feature set, including transaction count, sent and received transaction counts, unique counterparties, sent and received amounts, maximum value, active days, and tokens used. This difference is important because the paper's robustness claim depends on the two feature spaces being related but not redundant.",
    )

    doc.add_heading("KMeans6 Behaviour Roles", level=2)
    cluster_sizes["share"] = cluster_sizes["wallet_count"] / cluster_sizes["wallet_count"].sum()
    role_table = cluster_sizes.merge(behavior[["cluster", "behavior_interpretation"]], on="cluster", how="left")
    role_table = role_table.rename(columns={"cluster": "Cluster", "wallet_count": "Wallets", "share": "Share", "behavior_interpretation": "Interpreted role"})
    add_dataframe_table(doc, role_table[["Cluster", "Wallets", "Share", "Interpreted role"]], max_rows=10)
    add_figure(doc, "dataverse_files/network_science_kmeans6_full/cluster_feature_profile_heatmap.png", "Figure 1. Feature-profile heatmap showing how the six KMeans roles differ across wallet-level behavioural features.", width=6.4)

    doc.add_heading("Key Finding", level=2)
    add_para(
        doc,
        "The six clusters form a practical vocabulary for describing wallet populations: low-activity wallets, activity-leaning wallets, high-intensity or persistent wallets, bridge-like high-activity states, receiver or flow-size-leaning wallets, and smaller large-value volatile profiles. These roles are not identity claims about users; they are behavioural descriptions derived from transaction records.",
    )

    doc.add_heading("3. Advisor Gap 1: Formal Cluster Correspondence", level=1)
    add_para(
        doc,
        "The advisor report identified that the cluster correspondence claim needed a reproducible definition. The project already contained Hungarian matching and four-distance comparisons, but the rule had to be made explicit for paper use.",
    )
    add_para(
        doc,
        "The finalized rule defines correspondence in two layers. First, cluster centroids are compared under each distance metric. Second, maximum bipartite matching is used so each source cluster maps to at most one target cluster. A pair is treated as a core correspondence when it appears under at least three of four metrics: cosine, weighted Euclidean, Mahalanobis, and Wasserstein.",
    )
    add_dataframe_table(doc, distance_consensus, max_rows=10)
    add_para(
        doc,
        "Five cluster pairs meet the core correspondence criterion. Two pairs are especially strong because they appear under all four metrics. This supports the advisor's original claim that several behavioural archetypes reappear across analytical settings.",
    )

    doc.add_heading("4. Advisor Gap 2: Statistical Test of Correspondence", level=1)
    add_para(
        doc,
        "The advisor report specifically asked whether the cross-dataset correspondence could be expected by chance. The new analysis tested this using a random-label permutation design over the overlapping wallet labels.",
    )
    corr_summary = corr_perm.copy()
    corr_summary = corr_summary[
        [
            "n_wallets",
            "observed_max_bipartite_overlap",
            "adjusted_rand_index",
            "normalized_mutual_information",
            "null_mean",
            "null_95pct_lower",
            "null_95pct_upper",
            "z_score",
            "empirical_right_tail_p_value",
        ]
    ]
    add_dataframe_table(doc, corr_summary, max_rows=5)
    add_figure(doc, "paper_process_outputs/cross_dataset_correspondence_permutation_null.png", "Figure 2. Cross-dataset correspondence permutation test. The observed optimal overlap is far above the randomized-label null distribution.", width=5.9)
    add_para(
        doc,
        "The observed maximum bipartite wallet-label overlap was 0.395, compared with a null mean of 0.254. The empirical right-tail p-value was 0.000999, meaning the observed alignment exceeded every random permutation under add-one smoothing.",
    )
    add_para(
        doc,
        "Important interpretation: this wallet-level overlap statistic is not the same object as the earlier distance-metric correspondence claim. The distance-metric result says cluster profiles correspond across analyses; the wallet-overlap test says direct address-level label alignment is significantly above chance but numerically more moderate. The paper should report both, but should not call the wallet-level overlap 75%.",
    )

    doc.add_heading("5. Advisor Gap 3: Graph-Based Independent Validation", level=1)
    add_para(
        doc,
        "Graph validation was already one of the strongest completed parts of the project. The feature-similarity graph tests whether wallets close in behavioural feature space also form same-cluster communities under KMeans6 labels.",
    )
    graph_table = pd.DataFrame(
        [
            {
                "Nodes": int(feature_graph.iloc[0]["nodes"]),
                "Edges": int(feature_graph.iloc[0]["edges"]),
                "Intra-cluster edge ratio": feature_graph.iloc[0]["intra_cluster_edge_ratio"],
                "Modularity": feature_graph.iloc[0]["cluster_label_modularity"],
                "Assortativity": feature_graph.iloc[0]["cluster_assortativity"],
                "ARI detected vs KMeans6": feature_graph.iloc[0]["ari_detected_communities_vs_kmeans6"],
            }
        ]
    )
    add_dataframe_table(doc, graph_table, max_rows=3)
    add_para(
        doc,
        "Modularity asks whether the graph naturally separates into communities aligned with the cluster labels. Assortativity asks whether wallets with the same behavioural label tend to connect to each other. High values for both support the conclusion that the clusters are not just mathematical output; they reflect structure in the behavioural similarity graph.",
    )
    perm_table = graph_perm[
        [
            "observed_intra_ratio",
            "random_mixing_expected_ratio",
            "same_cluster_lift",
            "permutation_null_mean",
            "null_95pct_lower",
            "null_95pct_upper",
            "permutation_p_value_right_tail",
        ]
    ]
    add_dataframe_table(doc, perm_table, max_rows=3)
    add_para(
        doc,
        "The random-label graph test showed an observed same-cluster edge ratio of 0.943, compared with a null mean of 0.167. Wallets were therefore about 5.66 times more likely to connect to wallets in their own KMeans6 cluster than expected under random label mixing.",
    )

    doc.add_heading("6. Advisor Gap 4: Feature-Set Independence", level=1)
    add_para(
        doc,
        "The advisor report warned that convergence across two feature sets is only meaningful if the feature sets are not simply duplicates of each other. The new Spearman correlation analysis tested this at the individual-wallet level using sampled overlapping wallets.",
    )
    add_dataframe_table(doc, feature_corr, max_rows=3)
    add_figure(doc, "paper_process_outputs/cross_feature_set_spearman_heatmap.png", "Figure 3. Spearman correlation heatmap between original and expanded feature sets. The feature spaces are related but not redundant.", width=6.4)
    add_para(
        doc,
        "The median absolute Spearman correlation was 0.261 and the mean absolute correlation was 0.322. Only 8.1% of feature pairs had absolute correlation at or above 0.8. This supports a moderate version of the robustness claim: the two feature sets overlap conceptually, but most feature pairs are not highly redundant.",
    )

    doc.add_heading("7. Advisor Gap 5: Individual-Level Temporal Stability", level=1)
    add_para(
        doc,
        "The advisor report distinguished aggregate distributional stability from individual-level temporal stability. The project already had monthly state assignments, and the new report packaging converted those outputs into paper-ready retention summaries.",
    )
    add_dataframe_table(doc, temporal, max_rows=3)
    state_display = state_retention.rename(columns={"from_state": "State", "self_transition_probability": "Self-transition probability", "transition_entropy_bits": "Entropy", "top_destinations": "Top destinations"})
    add_dataframe_table(doc, state_display, max_rows=10)
    add_figure(doc, "paper_process_outputs/paper_ready_monthly_retention_rates.png", "Figure 4. Month-to-month individual wallet retention rates across 11 consecutive monthly transitions.", width=6.2)
    add_figure(doc, "dataverse_files/cluster_transition_analysis_expanded/figures/overall_transition_probability_matrix.png", "Figure 5. Overall transition probability matrix across monthly behavioural states.", width=5.5)
    add_para(
        doc,
        "Across 8,130,751 consecutive wallet-month pairs, the mean self-transition rate was 0.454, with monthly values ranging from 0.432 to 0.480. This means monthly behaviour is moderately persistent. The states are not fixed identities, but they are also not random monthly labels.",
    )
    add_para(
        doc,
        "C3 was the most stable state, with a self-transition probability of 0.524. C2 was also stable, at 0.519. C5 was the least stable, at 0.234. These results support the interpretation that some behavioural states are persistent roles while others are more transitional.",
    )

    doc.add_heading("8. Advisor Gap 6: Minimum Observation Experiment", level=1)
    add_para(
        doc,
        "The advisor report asked how much observation density is needed to recover core behavioural roles. The new experiment subsampled the sparse interval feature table at 10%, 25%, 50%, 75%, and 100%, reran KMeans on each retained fraction, and measured correspondence to the existing reference labels.",
    )
    add_dataframe_table(doc, obs, max_rows=10)
    add_figure(doc, "paper_process_outputs/minimum_observation_correspondence_curve.png", "Figure 6. Minimum observation experiment showing label-recovery correspondence across retained sparse-data fractions.", width=5.9)
    add_para(
        doc,
        "The experiment suggests that some correspondence can be recovered even at lower observation fractions. Correspondence to cluster_2 was high at 10% and 25%, while correspondence to cluster_1 peaked around 50% in this run. This should be framed as a practical subsampling recovery analysis rather than a definitive universal threshold.",
    )

    doc.add_heading("9. Alternative k, Silhouette, and Cluster Separability", level=1)
    add_para(
        doc,
        "The project also tested whether other cluster counts might be plausible. This does not invalidate KMeans6; instead, it shows the tradeoff between coarse segmentation and finer behavioural interpretation.",
    )
    alt_display = alt_k.rename(columns={"k": "k", "silhouette": "Silhouette", "ari_vs_current_kmeans6": "ARI vs KMeans6", "inertia": "Inertia"})
    add_dataframe_table(doc, alt_display, max_rows=10)
    add_para(
        doc,
        "The alternative-k results show that k=5 can produce a higher silhouette than k=6 in the full feature setting. KMeans6 remains defensible because the sixth cluster adds interpretability for narrower behaviour profiles and because the graph/statistical validation supports meaningful separation. The manuscript should acknowledge this openly: KMeans6 is not the only possible segmentation, but it is a useful and validated behavioural resolution.",
    )

    doc.add_heading("10. Transaction Network and Role-to-Role Interactions", level=1)
    add_para(
        doc,
        "Beyond feature similarity, the project examined how roles interact in transaction space. The role adjacency heatmaps and cluster-to-cluster edge dynamics show that role interactions are uneven: some roles interact heavily within themselves, while others serve as connectors across the behavioural state space.",
    )
    add_figure(doc, "dataverse_files/network_science_kmeans6_full/mean_monthly_cluster_adjacency_heatmap.png", "Figure 7. Mean monthly cluster adjacency heatmap showing role-to-role transaction interaction intensity.", width=5.8)
    add_para(
        doc,
        "This evidence is useful for interpretation because it connects feature-defined roles to actual transaction flows. A wallet role is therefore not only a location in feature space; it also has a characteristic interaction pattern with other roles.",
    )

    doc.add_heading("11. What We Have Found Overall", level=1)
    findings = [
        ("Finding 1", "Wallet behaviour can be compressed into interpretable roles.", "The KMeans6 model converts millions of addresses into six behavioural profiles that can be discussed without inspecting each wallet individually."),
        ("Finding 2", "The roles are statistically separable.", "Global and pairwise non-parametric tests show strong feature differences across clusters, with large effect sizes in key behavioural dimensions."),
        ("Finding 3", "The roles align with graph structure.", "Feature-similarity graph validation shows high same-cluster edge concentration, high modularity, high assortativity, and a strong random-label rejection."),
        ("Finding 4", "Cross-analysis correspondence exists but must be phrased carefully.", "Distance-metric profile matching shows several stable correspondences; direct wallet-label overlap is significant above chance but more moderate."),
        ("Finding 5", "The feature sets are related but not redundant.", "Cross-feature Spearman analysis shows moderate average association, supporting the claim that convergence is not merely caused by duplicate features."),
        ("Finding 6", "Monthly behaviour has moderate persistence.", "Individual wallet states show 43-48% month-to-month self-transition rates, with C2 and C3 among the most stable states."),
        ("Finding 7", "Some roles appear to require denser observation.", "The minimum observation experiment and cluster-profile differences suggest that finer-grained profiles are more sensitive to observation density than broad core roles."),
    ]
    findings_df = pd.DataFrame(findings, columns=["Finding", "Claim", "Evidence"])
    add_dataframe_table(doc, findings_df, max_rows=10, font_size=8.2)

    doc.add_heading("12. Paper Structure Following Advisor Guidance", level=1)
    add_para(
        doc,
        "The advisor HTML recommended a paper structure that begins with the research problem, introduces stability literature, defines both datasets and feature sets, formalizes correspondence, then presents results without overclaiming. Based on the work now completed, the paper can be organized as follows.",
    )
    structure = pd.DataFrame(
        [
            ["Introduction", "Problem, research questions, contribution statement", "Use the thesis that on-chain behavioural roles may be stable structural properties."],
            ["Background", "On-chain behavioural analytics and cluster stability literature", "Still needs literature writing; use Ben-David, von Luxburg, and cluster validity literature."],
            ["Data and Features", "Full-year ERC-20 wallet features and sparse interval features", "Define all feature groups and explain why the two feature sets differ."],
            ["Methods", "KMeans, correspondence metric, distance metrics, graph validation, temporal states, observation-density design", "Now supported by reproducible scripts and output tables."],
            ["Results", "Cluster profiles, correspondence, graph validation, temporal stability, feature correlation, minimum observation", "Use the figures and tables generated in this report."],
            ["Discussion", "Why roles persist, implications for simulation grounding and wallet analytics, limits", "Keep claims empirical; do not imply identity or intent."],
            ["Conclusion", "Answer research questions directly", "Emphasize validated behavioural segmentation and remaining external-replication opportunities."],
        ],
        columns=["Section", "Purpose", "How to write it now"],
    )
    add_dataframe_table(doc, structure, max_rows=10, font_size=8.2)

    doc.add_heading("13. Remaining Work and Limitations", level=1)
    add_para(
        doc,
        "The advisor's non-negotiable analytical list is now largely covered. The remaining work is mostly manuscript assembly, literature positioning, and careful interpretation.",
    )
    add_bullets(
        doc,
        [
            "Write the cluster stability literature section and position the project as empirical validation rather than a new clustering algorithm.",
            "Separate profile-level distance correspondence from wallet-level label-overlap correspondence; they answer related but different questions.",
            "Avoid overclaiming the 75% figure unless it is explicitly defined as a distance-metric profile correspondence result.",
            "Frame the minimum observation experiment as a subsampling recovery analysis, not a universal threshold for all chains or time periods.",
            "State limitations clearly: ERC-20 focus, no UTXO-chain replication, possible partial circularity in feature-graph validation, and no external ground-truth labels for behavioural intent.",
            "Treat bot or burst-activity interpretations as hypotheses unless external bot-label evidence is added.",
        ],
    )

    doc.add_heading("Final Assessment", level=1)
    add_para(
        doc,
        "The project has moved from a broad AI-governed simulation idea into a coherent empirical study of Ethereum wallet behaviour. The completed analyses now cover the advisor's main methodological concerns: formal correspondence, significance testing, individual temporal stability, feature-set non-redundancy, graph validation, and observation-density sensitivity.",
    )
    add_para(
        doc,
        "The most important conclusion is not that every wallet keeps the same exact label across every analysis. The stronger and more defensible conclusion is that Ethereum wallet populations contain recurring behavioural structures that can be recovered, interpreted, and validated using multiple independent forms of evidence.",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_report())
