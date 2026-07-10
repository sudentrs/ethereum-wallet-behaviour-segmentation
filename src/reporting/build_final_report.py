from pathlib import Path
import csv

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "ENS491-492_Final_Report_Ethereum_Wallet_Segmentation.docx"


def read_csv(path):
    with open(ROOT / path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if bold:
        run.font.color.rgb = RGBColor(31, 78, 121)


def style_table(table, header_fill="EAF2F8"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9)
            if i == 0:
                set_cell_shading(cell, header_fill)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, size=9)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, size=8.5)
    style_table(table)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)


def add_figure(doc, rel_path, caption, width=5.8):
    path = ROOT / rel_path
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(11)


def configure_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 20, RGBColor(31, 78, 121)),
        ("Heading 1", 15, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(47, 84, 150)),
        ("Heading 3", 11, RGBColor(31, 78, 121)),
    ]:
        st = styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color

    header = sec.header.paragraphs[0]
    header.text = "ENS 491-492 Graduation Project Final Report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Ethereum Wallet Behaviour Segmentation and Network Science Analysis")
    footer.runs[0].font.name = "Arial"
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)


def build():
    cluster_rows = read_csv("dataverse_files/network_science_kmeans6_full/cluster_behavior_interpretation.csv")
    summary = read_csv("dataverse_files/network_science_kmeans6_full/feature_similarity_network_summary.csv")[0]
    perm = read_csv("dataverse_files/network_science_kmeans6_full/feature_graph_random_label_hypothesis_test.csv")[0]
    altk = read_csv("dataverse_files/network_science_kmeans6_full/alternative_k_clusterability.csv")
    coverage = read_csv("dataverse_files/network_science_kmeans6_full/monthly_transaction_network_coverage.csv")

    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ENS 491-492\nGraduation Project\nFinal Report")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(31, 78, 121)
    doc.add_paragraph()
    meta = [
        ("Project Title", "Ethereum Wallet Behaviour Segmentation and Network Science Analysis for AI-Governed Crypto Finance Simulation"),
        ("Group Members", "Sena Kahya, Sude Tiras"),
        ("Supervisor(s)", "Selim Balcisoy"),
        ("Date", "May 2026"),
    ]
    add_table(doc, ["Field", "Information"], meta, widths=[1.6, 5.4])
    add_para(doc, "This report follows the ENS 491-492 final report template. The project evolved from an initial LLM-governed crypto-bank simulator concept into a data-driven Ethereum wallet behaviour segmentation and network-science validation study. The final implementation focuses on deriving empirical wallet roles from blockchain transaction behaviour, validating those roles statistically and graph-theoretically, and preparing the behavioural state space needed for later explainable simulation and advisory experiments.")
    doc.add_page_break()

    add_heading(doc, "EXECUTIVE SUMMARY", 1)
    add_para(doc, "This project researched how Ethereum wallet behaviour can be segmented into interpretable user types and validated as meaningful network communities. The original motivation was to support an AI-governed crypto finance simulator with realistic synthetic agents, but the implementation showed that credible agent initialization first requires empirical behavioural groups rather than manually invented personas. The final work therefore developed a data-driven pipeline that processes Ethereum transaction records, extracts wallet-level behavioural features, applies robust scaling and KMeans clustering, and evaluates the resulting six wallet segments through network science, statistical testing, transaction-flow analysis, and temporal transition modelling.")
    add_para(doc, "The core methodology combines unsupervised learning with graph-theoretic validation. Wallets are represented by activity, value-flow, directionality, persistence, volatility, and temporal features. KMeans6 assigns wallets to behavioural clusters, and a feature-similarity graph tests whether wallets with similar profiles connect preferentially to the same cluster. The result is strong evidence of non-random structure: the feature-similarity graph has 23,755 sampled nodes, 80,596 edges, 0.943 intra-cluster edge ratio, 0.777 modularity, and 0.932 assortativity. A random-label permutation test produced a z-score of 592.5 and p = 0.000999, showing that the observed cluster alignment is far beyond random mixing. Transaction-network analysis further shows that clusters occupy distinct roles, including central, bridge, receiver, low-activity, and niche high-value behaviours. Temporal analysis shows that full-year clusters act as long-term identities while monthly state assignments capture short-term behavioural expression. Overall, the project produced a scientifically defensible wallet segmentation framework and a foundation for future explainable crypto-finance simulation.")

    add_heading(doc, "PROBLEM STATEMENT", 1)
    add_heading(doc, "Original Problem and Motivation", 2)
    add_para(doc, "Decentralised finance and blockchain-based financial systems generate large amounts of transaction data, but user behaviour in these systems is difficult to summarize. Wallets are pseudonymous, transaction distributions are heavy-tailed, and apparent behaviour can change across time. A simulation or credit-scoring system that relies on simple hand-written personas such as optimistic, neutral, fearful, or speculative users risks producing unrealistic outputs because those personas are not anchored in actual transaction behaviour. The problem addressed in this project is therefore the construction and validation of empirically grounded crypto-wallet behaviour groups that can serve as the user layer of a broader AI-governed crypto finance simulation.")
    add_para(doc, "The motivation is both scientific and practical. Scientifically, blockchain data provides a rich setting for studying heavy-tailed activity, community structure, centrality, temporal transitions, and behavioural role formation. Practically, any AI-mediated financial advisory or credit-scoring simulator needs realistic agent types before it can responsibly test advisory effects, fairness, explainability, or risk outcomes. By replacing manually assigned personas with data-driven wallet segments, the project improves the credibility of downstream simulation results and provides a clearer explanation of how agents are initialized.")
    add_heading(doc, "Relation to Prior Work", 2)
    add_para(doc, "Prior work in agent-based financial simulation, such as ABIDES, demonstrates how complex market behaviour can be studied in controlled synthetic environments. Work on DeFi systems such as MakerDAO shows that algorithmic lending and collateral mechanisms can be implemented at scale. Explainable credit-scoring literature highlights the need for interpretability when automated models affect financial decisions. Insider-trading and market-surveillance studies also show that unsupervised methods can identify unusual or coordinated trading behaviour from large financial datasets. However, these strands do not fully solve the present problem: a crypto-finance simulator needs empirically derived wallet behaviour types, validation that those types are not arbitrary, and a temporal state model that explains how wallet roles change over time.")
    add_para(doc, "This project contributes to that gap by using Ethereum transaction data to create wallet behaviour segments, validating the segmentation with non-parametric statistics and network science, and interpreting the clusters as behavioural roles that can later be used in simulation, advisory-response analysis, and explainable scoring.")

    add_heading(doc, "Objectives/Tasks", 2)
    add_table(doc, ["Objective / Task", "Intended Result", "Final Status"], [
        ("Process Ethereum transaction data", "Create wallet-level datasets suitable for behavioural analysis.", "Completed through monthly transaction processing and wallet feature construction."),
        ("Engineer behavioural features", "Represent activity intensity, persistence, value scale, flow direction, volatility, and timing.", "Completed with expanded robust-scaled feature sets."),
        ("Segment wallets into behavioural groups", "Identify natural user types for agent initialization.", "Completed using KMeans6 as the primary interpretable segmentation."),
        ("Validate the segmentation", "Show clusters are statistically separable and graph-theoretically meaningful.", "Completed through Kruskal-Wallis, Mann-Whitney, Cliff's delta, modularity, assortativity, conductance, and permutation tests."),
        ("Analyse transaction-network roles", "Interpret clusters as system-level roles in raw transfer networks.", "Completed for selected 2022 months with directed weighted cluster graphs."),
        ("Analyse temporal behaviour", "Distinguish long-term identity from monthly behavioural state.", "Completed using a Markov-style transition matrix over monthly states."),
        ("Prepare for explainable simulation", "Provide empirical agent groups and features for future advisory and credit-score experiments.", "Partially completed: segmentation and validation are ready; full LLM-advisory simulation remains future work."),
    ], widths=[2.0, 2.8, 2.2])

    add_heading(doc, "Realistic Constraints", 2)
    add_para(doc, "The project was affected by several realistic constraints. The first constraint was data volume: raw Ethereum transaction files are large, so the implementation required monthly processing, aggregation, sampling for graph construction, and memory-conscious scripts. The second constraint was interpretability. A model that only optimizes a clustering metric would not be sufficient for a graduation project; the clusters also needed behavioural explanations and statistical evidence. The third constraint was reproducibility. The pipeline relies on deterministic feature extraction, saved CSV outputs, and documented scripts rather than opaque manual inspection. The fourth constraint was privacy and ethics: wallet addresses are pseudonymous public blockchain identifiers, but the project avoids personal identification and reports aggregate behavioural groups rather than attempting to deanonymize individuals. Finally, computational and environmental constraints encouraged efficient offline analysis on standard hardware instead of costly continuous blockchain indexing or large-scale model training.")
    add_para(doc, "No formal engineering code such as ISO/IEC 25010 was mandatory for the experimental prototype, but software-quality principles from reproducibility, maintainability, and explainability guided the design. The work also aligns with responsible data-science practices by avoiding real credit decisions, financial recommendations, or attempts to classify named persons.")
    add_para(doc, "Economic constraints were also relevant. A fully indexed blockchain analytics platform would normally require database infrastructure, continuous ingestion, and large storage capacity. The project instead used a research pipeline based on monthly files, local feature extraction, and saved intermediate tables. This made the project feasible on standard hardware and kept the computational budget compatible with a student project. The main trade-off is that the system is optimized for offline analysis rather than real-time monitoring.")
    add_para(doc, "Environmental constraints were addressed indirectly through the same design choices. The solution does not execute smart contracts, mine blocks, or add activity to the Ethereum network. Its computational footprint comes from reading, cleaning, aggregating, and modelling existing transaction data. Sampling was used for graph construction because a full pairwise wallet-similarity graph over millions of wallets would be computationally expensive and unnecessary for validation. This sampling strategy preserves the scientific question while reducing runtime, memory use, and energy consumption.")
    add_para(doc, "Health and safety constraints are not physical in this project, but financial safety matters. Because the work concerns wallet behaviour and possible future credit scoring, the prototype avoids user-facing recommendations. No wallet owner is contacted, scored for a real service, or labelled as risky in a way that could affect access to finance. The project remains an analytical framework and a simulation foundation.")

    doc.add_page_break()
    add_heading(doc, "METHODOLOGY", 1)
    add_heading(doc, "Data Sources and Processing Pipeline", 2)
    add_para(doc, "The project uses Ethereum transaction data organized by month. The raw transaction files were transformed into wallet-level feature datasets. The processing pipeline reads monthly transactions, identifies sender and receiver behaviour, aggregates activity by wallet, and creates features that summarize both total-year and monthly behaviour. The final KMeans6 cluster file used as the authoritative segmentation is dataverse_files/wallet_kmeans_clusters.csv. Additional expanded feature sets and monthly wallet files were produced for robustness, subcluster, and temporal analyses.")
    add_para(doc, "The pipeline is modular. Separate scripts process raw data, extract features, scale values, perform clustering, analyse transaction networks, test cluster separation, and generate figures. This modular structure was important because the project evolved from a simulation-first plan to a segmentation-first plan. Once the segmentation stage became central, each later analysis could reuse the same cluster labels and feature definitions.")
    add_para(doc, "The first stage of the pipeline is raw data preparation. Monthly transaction files are read and reduced to the fields needed for behavioural analysis, such as sender address, recipient address, token value or transaction amount, and date. Because the raw files are large, the processing is organized by month instead of loading the full dataset at once. Intermediate outputs are saved as monthly wallet tables. This structure makes the pipeline recoverable: if one month needs to be corrected, the whole year does not need to be recomputed.")
    add_para(doc, "The second stage is wallet aggregation. Each wallet can appear as a sender, recipient, or both. The feature extraction scripts therefore compute both outgoing and incoming activity, then combine them into wallet-level behavioural summaries. This is important because wallets can play different roles in a crypto ecosystem. Some wallets primarily receive value, some primarily send value, some operate as high-throughput hubs, and many appear only briefly. A sender-only view would miss receiver and sink behaviours, while a receiver-only view would miss active dispatcher behaviour.")
    add_para(doc, "The third stage is feature normalization. Ethereum activity is not normally distributed. Most wallets have low or moderate activity, while a small number have extremely large transaction counts, values, or flow volatility. The pipeline therefore uses transformations and robust scaling so that distance-based clustering reflects behavioural patterns rather than being dominated by a few extreme numeric values. Robust scaling is especially appropriate because it uses median and interquartile information, making it less sensitive to outliers than standard mean-variance scaling.")
    add_para(doc, "The fourth stage is model production. Clustering outputs are saved in CSV form and treated as explicit research artifacts rather than temporary notebook variables. This allows later scripts to perform independent validation. For example, the network-science scripts can load the same cluster assignments and ask whether they form communities in a wallet-similarity graph. The temporal scripts can load the same prototypes and ask whether monthly wallet behaviour moves between states. This artifact-based design improves reproducibility and makes the final report traceable to concrete files.")
    add_heading(doc, "Feature Engineering", 2)
    add_para(doc, "Wallet behaviour was represented through features that capture multiple aspects of financial activity. Activity features include total transactions, average transactions per day, active days, and activity density. Temporal features include activity span, burstiness, and transaction-size volatility. Value-flow features include average transaction size, net flow, and flow volatility. Directionality features include in-out ratio and directional imbalance between receiving and sending. Expanded daily features also track monthly or daily totals, maximum values, standard deviations, concentration ratios, and active-day patterns.")
    add_para(doc, "Because blockchain activity is heavy-tailed, raw values were transformed and robust-scaled before clustering. The descriptive diagnostics support this decision: transaction counts and value flows have extreme p99-to-median and maximum-to-median ratios. For example, the median total transaction count in a sample was 20, while the p99 value was 956 and the maximum exceeded 2.6 million. Such distributions would cause ordinary distance-based clustering to be dominated by extreme wallets if not scaled carefully.")
    add_para(doc, "The feature set was designed to separate what a wallet does from how often it does it and when it does it. Total transaction count alone cannot distinguish a stable daily user from a wallet that performs the same number of transactions in a short burst. Similarly, total flow alone cannot distinguish a high-value receiver from a high-value sender. For this reason, the project uses multiple feature families: intensity, persistence, value scale, directionality, concentration, and volatility. These families make the resulting clusters easier to interpret and less dependent on a single metric.")
    add_para(doc, "Activity intensity measures how much a wallet participates in the transaction ecosystem. This includes total transaction count, average transactions per active period, and activity density. Persistence measures whether behaviour is spread across time or concentrated in a short window. Value-scale features measure the magnitude of transfers, while flow features measure net incoming or outgoing tendency. Directionality captures whether a wallet is balanced or strongly one-sided in receiving versus sending. Volatility features measure irregularity in transaction size or daily activity.")
    add_para(doc, "Feature engineering also supports the future simulation goal. In an agent-based simulator, agent types need parameters that affect behaviour. The features produced here can become those parameters: high activity intensity can translate into more frequent simulated decisions; strong directionality can translate into receiver- or sender-biased behaviour; volatility can translate into larger state changes under market events; and persistence can influence how likely an agent is to remain active across simulation steps. Thus, the feature design is not only descriptive; it also prepares a bridge from real wallet data to synthetic agent behaviour.")
    add_heading(doc, "Clustering Approach", 2)
    add_para(doc, "The project explored unsupervised learning methods, including centroid-based KMeans and density-oriented alternatives such as HDBSCAN. The final report focuses on KMeans6 because it produced a compact, interpretable six-role segmentation and because subsequent network-science analysis showed strong community alignment. KMeans partitions wallets into six clusters in robust-scaled feature space. The choice of six clusters is not claimed to be mathematically optimal for every possible objective; rather, it is defended as a useful behavioural resolution that preserves important small and high-value groups while remaining explainable.")
    add_table(doc, ["k", "Silhouette", "ARI vs KMeans6", "Inertia"], [
        (r["k"], f'{float(r["silhouette"]):.3f}', f'{float(r["ari_vs_current_kmeans6"]):.3f}', f'{float(r["inertia"]):,.0f}') for r in altk
    ], widths=[0.8, 1.5, 1.6, 1.8])
    add_para(doc, "The alternative-k diagnostics show that k=5 has the highest silhouette among the tested values, while k=6 remains defensible because it maintains a small but highly distinctive high-value directional cluster and supports richer role interpretation. This is a design trade-off between compactness and behavioural specificity.")
    add_para(doc, "KMeans was selected for the primary segmentation because it creates explicit cluster centroids or prototypes. This is useful for interpretation and for later temporal analysis. A centroid-like prototype can be used as a reference point: monthly wallet states can be assigned to the nearest full-year prototype, and synthetic agents can be initialized around cluster-level behaviour profiles. Density-based methods are useful for detecting irregular groups and outliers, but they can be harder to convert into a simple state space for simulation.")
    add_para(doc, "The project does not treat the KMeans objective as the only proof of validity. This is an important methodological decision. KMeans will always partition data into k groups, even if the groups are not meaningful. The report therefore validates the labels after clustering through independent evidence: graph community structure, statistical separation, effect sizes, transaction-network roles, temporal transition behaviour, and internal subcluster analysis. This reduces the risk of over-interpreting an algorithmic artifact.")
    add_para(doc, "The cluster labels are interpreted as behavioural segments rather than identities of real-world persons. A wallet in C1 is not claimed to be a specific type of human or institution. Instead, C1 means that the wallet's observed transaction behaviour has high activity intensity, persistence, and value-flow characteristics relative to the feature set. This distinction matters ethically and scientifically because blockchain addresses may be controlled by individuals, exchanges, contracts, bots, or groups of users.")
    add_heading(doc, "Feature-Similarity Network Construction", 2)
    add_para(doc, "To test whether KMeans6 labels are meaningful beyond a distance objective, a feature-similarity graph was constructed. In this graph, nodes are sampled wallets and edges connect wallets that are mutual nearest neighbours in robust-scaled feature space. The graph therefore represents behavioural similarity rather than direct transaction flow. If KMeans labels are meaningful, same-cluster wallets should be more densely connected to one another than expected by random label assignment.")
    add_table(doc, ["Statistic", "Value"], [
        ("Nodes", f'{int(summary["nodes"]):,}'),
        ("Edges", f'{int(summary["edges"]):,}'),
        ("Density", f'{float(summary["density"]):.6f}'),
        ("Intra-cluster edge ratio", f'{float(summary["intra_cluster_edge_ratio"]):.3f}'),
        ("Inter-cluster edge ratio", f'{float(summary["inter_cluster_edge_ratio"]):.3f}'),
        ("KMeans6 label modularity", f'{float(summary["cluster_label_modularity"]):.3f}'),
        ("Cluster assortativity", f'{float(summary["cluster_assortativity"]):.3f}'),
        ("Detected communities", summary["detected_communities"]),
    ], widths=[3.2, 2.0])
    add_figure(doc, "dataverse_files/network_science_kmeans6_full/cluster_feature_profile_heatmap.png", "Figure 1. Robust-scaled feature profile heatmap for KMeans6 clusters.", width=5.8)
    add_para(doc, "Several network measures are used because each answers a different question. The intra-cluster edge ratio asks how often similarity edges stay inside the same KMeans label. Modularity asks whether the partition creates more within-cluster edges than expected under a degree-preserving null model. Assortativity asks whether adjacent nodes tend to share the same categorical label. Conductance asks how strongly a cluster is separated from the rest of the graph. Together, these metrics provide a richer validation than any single score.")
    add_para(doc, "The distinction between the feature-similarity graph and the transaction graph is central. In the feature-similarity graph, an edge means two wallets behave similarly according to engineered features. In the transaction graph, an edge means value or transaction count moved from one behavioural group to another. A segmentation can be meaningful in one graph and not the other. The project intentionally analyses both because a useful wallet segmentation should be coherent in behaviour space and interpretable in transaction-flow space.")
    add_heading(doc, "Statistical Validation", 2)
    add_para(doc, "The project uses non-parametric statistics because wallet features are non-normal, skewed, and heavy-tailed. Global Kruskal-Wallis tests evaluate whether each feature differs across the six clusters. Pairwise Mann-Whitney tests and Cliff's delta quantify differences between cluster pairs. Benjamini-Hochberg false-discovery-rate correction is used for multiple comparisons. This combination allows the report to distinguish statistical significance from practical effect size.")
    add_para(doc, "A random-label permutation test provides the strongest graph-level validation. The null hypothesis is that KMeans6 labels are independent of the feature-similarity graph. Labels are randomly permuted while preserving cluster sizes, and the intra-cluster edge ratio is recomputed for each permutation. The observed same-cluster edge ratio is then compared with the random-label distribution.")
    add_table(doc, ["Permutation-Test Quantity", "Value"], [
        ("Observed intra-cluster ratio", f'{float(perm["observed_intra_ratio"]):.4f}'),
        ("Random-label expected ratio", f'{float(perm["random_mixing_expected_ratio"]):.4f}'),
        ("Same-cluster lift", f'{float(perm["same_cluster_lift"]):.2f}x'),
        ("Null 95% interval", f'[{float(perm["null_95pct_lower"]):.4f}, {float(perm["null_95pct_upper"]):.4f}]'),
        ("Permutation z-score", f'{float(perm["permutation_z_score"]):.2f}'),
        ("Empirical right-tail p-value", f'{float(perm["permutation_p_value_right_tail"]):.6f}'),
    ], widths=[3.2, 2.4])
    add_para(doc, "The global Kruskal-Wallis tests show that the six clusters differ across all major feature dimensions. The most globally separating features include in_out_ratio, directionality, average transaction size, net flow, flow volatility, active days, activity span, average transactions per day, activity density, total transactions, transaction-size volatility, and burstiness. This means the clusters are not differentiated by one isolated variable. They are separated across a broad behavioural profile.")
    add_para(doc, "Pairwise testing provides a more detailed view. Some cluster pairs are especially distinct, such as C1-C2, C3-C5, C3-C4, C2-C3, and C0-C1. These differences have strong mean absolute Cliff's delta values, indicating that the separation is practically meaningful, not only statistically significant. This matters because very large datasets can produce very small p-values for minor effects. Effect-size reporting helps avoid that problem.")
    add_para(doc, "The permutation test is the strongest non-parametric argument because it does not assume a feature distribution. It fixes the similarity graph and asks what would happen if the same cluster sizes were randomly assigned to graph nodes. The observed intra-cluster ratio is outside the random-label 95% interval by a very large margin. Since the random-label null mean is approximately 0.1666 and the observed value is approximately 0.9432, the result supports the claim that the clustering is aligned with actual structure in feature space.")
    add_heading(doc, "Transaction-Network Analysis", 2)
    add_para(doc, "A second graph representation was built from raw transactions. In this graph, the six clusters are nodes and directed weighted edges represent transfer counts from one cluster to another. This graph does not test feature similarity; it tests how behavioural groups interact in transaction flow. The analysis was performed for selected high-volume months in 2022. Sender and recipient cluster labels were available for roughly three quarters of raw rows, creating enough coverage to interpret cluster-level flows.")
    add_table(doc, ["Month", "Rows Read", "Rows with Both Labels", "Coverage", "Self-loop Ratio", "Reciprocity"], [
        (r["month"], f'{int(r["raw_rows_read"]):,}', f'{int(r["rows_with_both_cluster_labels"]):,}', f'{float(r["coverage_ratio"]):.3f}', f'{float(r["self_loop_ratio"]):.3f}', f'{float(r["reciprocity_binary"]):.3f}') for r in coverage
    ], widths=[0.8, 1.3, 1.5, 0.9, 1.1, 1.0])
    add_figure(doc, "dataverse_files/network_science_kmeans6_full/mean_monthly_cluster_adjacency_heatmap.png", "Figure 2. Mean monthly cluster-to-cluster transaction adjacency heatmap.", width=5.6)
    add_para(doc, "The transaction network uses a six-node directed weighted graph for each analysed month. Each node is one behavioural cluster. An edge from C_a to C_b counts transfers from wallets assigned to C_a to wallets assigned to C_b. Self-loops represent transfers where sender and recipient belong to the same behavioural cluster. This aggregation makes a very large transaction table interpretable at the behavioural-role level.")
    add_para(doc, "The high self-loop ratios in May and June 2022 suggest that wallets with similar long-term behavioural profiles often transact within the same group. The lower self-loop ratio in September 2022 suggests a period of more cross-cluster interaction. This may reflect market conditions, changes in active users, token-flow shocks, or other external events. The project does not claim a causal explanation for the September pattern, but it identifies a clear temporal change in cluster-level mixing.")
    add_para(doc, "Centrality measures identify role differences. C1 has very large weighted in-strength and out-strength, which makes it a central activity hub. C4 often acts as a receiver or sink because it has strong incoming external flow. C3 becomes important as a bridge-like or connector cluster, particularly when its flows to and from C1 increase. These transaction roles reinforce the behavioural interpretation derived from features.")
    add_heading(doc, "Temporal State Analysis", 2)
    add_para(doc, "The project also models monthly behavioural transitions. Full-year KMeans6 labels are interpreted as long-term behavioural identities. For each active wallet-month, a monthly feature vector is assigned to the nearest full-year cluster prototype. This produces a monthly state sequence for each wallet. State transitions are summarized as a Markov-style transition matrix, where each entry is the probability of moving from one monthly state to another between consecutive months.")
    add_para(doc, "This method separates stable identity from short-term expression. A wallet may belong to full-year cluster C1 but behave like C3 in a particular month. The resulting transition matrix identifies stable states, transitional states, and bridge states.")
    add_figure(doc, "dataverse_files/cluster_transition_analysis_expanded/figures/overall_transition_probability_matrix.png", "Figure 3. Overall monthly behavioural transition probability matrix.", width=5.7)
    add_figure(doc, "dataverse_files/cluster_transition_analysis_expanded/figures/monthly_behavior_state_composition.png", "Figure 4. Monthly composition of behavioural states.", width=5.7)
    add_para(doc, "The temporal model uses the full-year KMeans6 prototypes as behavioural reference points. For each active wallet-month, the monthly feature vector is assigned to the closest prototype. This creates a state sequence such as C2, C3, C3, C4 for a wallet that changes behaviour over time. The state sequence does not replace the full-year label. Instead, it adds a dynamic layer over the long-term identity.")
    add_para(doc, "The transition matrix shows that C2 and C3 are the most stable monthly states, with self-transition probabilities slightly above 0.51. C1 is moderately stable with a self-transition probability of about 0.457. C0, C4, and C5 are more transitional. This result is important for simulation: some agents should have persistent behavioural roles, while others should be more likely to move between states when conditions change.")
    add_para(doc, "The strongest off-diagonal pathways include C2 to C3, C3 to C2, C4 to C3, C5 to C3, C0 to C3, C1 to C0, C5 to C4, and C4 to C2. The repeated movement into C3 supports the interpretation of C3 as a bridge or attractor state. In a future simulator, C3 can represent a state that many agents enter during active or transitional periods.")
    add_heading(doc, "Internal Subcluster Analysis", 2)
    add_para(doc, "The final analytical layer asks whether each main KMeans6 cluster is internally homogeneous. For each main cluster, the project tested internal subcluster counts from 2 to 8 and selected interpretable internal partitions using silhouette and size constraints. Internal Kruskal-Wallis tests and pairwise effect-size calculations show that each main cluster contains meaningful substructure. This is especially strong for C5, whose internal silhouette is 0.726, and C3, whose internal silhouette is 0.364.")
    add_figure(doc, "dataverse_files/cluster_behavior_subclusters_expanded/figures/subcluster_silhouette_by_main_cluster.png", "Figure 5. Internal subcluster silhouette by main KMeans6 cluster.", width=5.6)
    add_para(doc, "Internal subcluster analysis is useful because a main cluster may be valid but still contain multiple subtypes. C0, for example, is interpreted as mixed or transitional, and internal analysis supports this by selecting three subclusters. C5 is small in the main segmentation but has very strong internal structure, suggesting that even the niche high-value directional group contains an extreme subtype and a broader large-value subtype. This type of result would be hidden if the project stopped at the main six labels.")
    add_para(doc, "The internal results also support a hierarchical interpretation. The six main clusters provide a manageable top-level behavioural taxonomy. Within that taxonomy, subclusters can provide more detail when needed. For simulation, this means the system could initialize broad agent classes from KMeans6 and then sample subtypes inside each class to create diversity. For reporting and explainability, the main clusters remain easier to communicate.")

    doc.add_page_break()
    add_heading(doc, "DISCUSSION", 1)
    add_heading(doc, "Realization of Initial Objectives", 2)
    add_para(doc, "The initial objective was to build a crypto-like financial simulation environment influenced by advisory AI, market events, and behaviour-based scoring. During implementation, the team identified a prerequisite: synthetic agents should not be initialized from arbitrary personas. They need empirical behavioural foundations. The final project therefore realized the most important data and modelling foundation for the simulator: a validated wallet segmentation and temporal state space. The LLM-advisory and SHAP-based simulation layer was not fully completed, but the segmentation, validation, and network interpretation components were developed more deeply than originally planned.")
    add_para(doc, "This change should be interpreted as a refinement rather than a failure of scope. A simulation built directly on predefined personas would have been easier to implement but weaker scientifically. By prioritizing data-driven segmentation, the project produced a more defensible base for future simulation, credit-score explainability, and advisory-impact experiments.")
    add_para(doc, "The initial objectives were therefore realized unevenly but productively. The data-processing, feature-engineering, segmentation, validation, and analysis objectives were realized in depth. The advisory LLM, dynamic credit scoring, SHAP dashboard, and full simulation-loop objectives were not fully realized as final implemented modules. However, the final work provides the empirical input those modules need. In other words, the project completed the foundation layer of the original system rather than a shallow version of every planned layer.")
    add_para(doc, "The difference from the initial objectives is most visible in the role of personas. The original plan treated personas as inputs to the simulator. The final plan treats behavioural segments as outputs of data analysis and future inputs to the simulator. This reversal improved the scientific quality of the project because it allowed the agent layer to emerge from observed financial behaviour.")
    add_heading(doc, "Behavioural Interpretation of Clusters", 2)
    compact_rows = []
    for r in cluster_rows:
        compact_rows.append((
            f'C{r["cluster"]}',
            f'{int(r["wallet_count"]):,}',
            r["behavior_interpretation"],
            r["top_high_features"],
        ))
    add_table(doc, ["Cluster", "Wallets", "Behavioural Interpretation", "Top High Features"], compact_rows, widths=[0.7, 1.0, 2.5, 2.8])
    add_para(doc, "The clusters represent distinct wallet roles. C1 is the high-intensity persistent activity cluster and also the main center in the transaction network. C2 is a low-activity, short-lived, low-volatility behaviour type. C3 is a dense activity and bridge-like state that becomes important in transaction flows and temporal transitions. C4 and C5 represent large-value and directional receiver/sender imbalance, with C5 being small but highly distinctive. C0 is more mixed and transitional, with higher conductance than the most separated clusters.")
    add_para(doc, "C0 can be interpreted as a small-value and net-outflow group with mixed characteristics. It does not have the strongest activity intensity, value scale, or directionality, but it has enough boundary edges to interact with other behavioural regions. This makes it a plausible transitional or baseline active-wallet segment. In later simulation work, C0 could represent ordinary wallets that shift behaviour more easily under external events.")
    add_para(doc, "C1 is the most central high-activity segment. It combines high transaction intensity, persistent activity, and larger value-flow behaviour. In the transaction graph it dominates weighted flow and PageRank-like centrality. This makes C1 a natural core actor in the ecosystem. In simulation, C1-like agents would likely have frequent decisions, strong exposure to market events, and large influence on aggregate system behaviour.")
    add_para(doc, "C2 is the low-activity and short-lived segment. It is large in wallet count but behaviourally quieter. This cluster is important because blockchain ecosystems contain many wallets that appear briefly or transact lightly. Treating these wallets as noise would miss a major part of the system. The temporal model also shows that C2 is one of the more stable monthly states, suggesting that low-activity behaviour can be persistent rather than merely incomplete observation.")
    add_para(doc, "C3 has one of the most interesting roles. Its feature profile suggests intense short-window activity, and its temporal profile suggests a bridge or attractor state. Many states transition into C3, and C3 connects strongly with C1 in transaction flows. This makes C3 a key candidate for modelling reactive behaviour under changing market conditions. If the future simulator includes shocks, C3-like behaviour may represent wallets that become active connectors during those shocks.")
    add_para(doc, "C4 and C5 capture large-value and directionally imbalanced behaviour. C4 is broader and more common, while C5 is much smaller and more extreme. C5's low conductance and high directionality make it especially distinctive. Preserving C5 is one reason KMeans6 remains useful even though k=5 has a slightly higher silhouette. A coarser model could merge away an important niche behaviour.")
    add_heading(doc, "Graph-Theoretic Findings", 2)
    add_para(doc, "The strongest result is that KMeans6 labels align with the topology of the wallet feature-similarity graph. The graph has a 0.943 intra-cluster edge ratio, meaning that most behavioural-neighbour edges connect wallets in the same cluster. The modularity value of 0.777 is high for a partition imposed on a graph, and assortativity of 0.932 indicates that graph edges strongly prefer same-label endpoints. The permutation test confirms that this is not an artifact of cluster sizes: random label assignments produce an expected intra-cluster ratio near 0.167, while the observed value is about 5.66 times higher.")
    add_para(doc, "Conductance adds cluster-level nuance. C5 has extremely low conductance, which supports the interpretation that it is a niche and highly separated behaviour type. C3 also has low conductance while playing an important dynamic role. C0 has the highest conductance among the main results, consistent with a more mixed or transitional role.")
    add_para(doc, "The graph-theoretic evidence matters because it validates the clustering using relationships among wallets rather than only feature averages. If clusters were merely numerical bins, the feature-similarity graph would not necessarily show strong same-cluster connectivity. The observed modularity and assortativity indicate that the clusters behave like communities in a network of behavioural similarity. This supports the use of network-science concepts from the course, including community structure, assortative mixing, conductance, small-world diagnostics, and robustness analysis.")
    add_para(doc, "The feature graph also creates a bridge between machine learning and network science. KMeans is a geometric method, but the graph converts geometric similarity into topology. Once topology is available, the project can ask network questions: Are clusters internally cohesive? Are boundaries weak or strong? Do labels align with communities? Could random labels produce the same pattern? The answer across these tests is consistently that KMeans6 has meaningful structure.")
    add_heading(doc, "Transaction-Network Findings", 2)
    add_para(doc, "The raw transaction graph shows that behavioural clusters also occupy different system-level roles. In the analysed months, C1 dominates transaction volume and acts as the main center. Self-loop ratios are high in May and June 2022, meaning many transfers occur within the same behavioural group. In September 2022, the self-loop ratio falls to 0.487, indicating more cross-cluster mixing. The C3-C1 and C1-C3 connections grow strongly in September, supporting the interpretation of C3 as a bridge or connector cluster under changing market conditions.")
    add_para(doc, "The transaction-network results should be interpreted carefully. Cluster-level aggregation hides individual wallet-level details, and the months analysed are selected windows rather than a complete causal study of all market events. Nevertheless, the method is valuable because it shows that clusters are not only separated in feature space. They also interact differently in actual transfer flow. This strengthens the behavioural interpretation of the clusters and makes them more useful for simulation.")
    add_para(doc, "For future work, transaction-network dynamics can become an event detector. If the self-loop ratio drops or a bridge cluster suddenly gains flow, the simulator could represent this as a market stress or behavioural contagion signal. The current project does not implement that feedback loop, but it identifies the necessary measurements.")
    add_heading(doc, "Temporal Findings", 2)
    add_para(doc, "Temporal analysis shows that wallet behaviour is only moderately persistent month to month. Monthly self-transition rates range from about 0.432 to 0.480, meaning that more than half of active wallets often shift to another behavioural state in the following month. C2 and C3 are the most stable monthly states, with self-transition probabilities around 0.519 and 0.524 respectively. C4 and C5 are more transitional, with lower self-transition probabilities. Several states move strongly into C3, including C2, C4, C5, and C0, which supports the interpretation of C3 as a dynamic attractor or bridge state.")
    add_para(doc, "The temporal result is one of the most important refinements to the original simulator idea. Personas are usually static, but wallet behaviour is not. A wallet can have a long-term identity and still express different monthly states. This supports a two-layer agent model: each agent has a long-term cluster identity, and at each simulation step the agent occupies a short-term state that may change according to transition probabilities and event effects.")
    add_para(doc, "This design also makes the future advisory experiment more realistic. An LLM advisory message may not change a wallet's permanent identity, but it may influence short-term state transitions. For example, under a volatile event, advisory tone might increase the probability that a C0-like agent moves into a C3-like active bridge state or that a C5-like high-value agent shifts toward a C4 receiver state. The transition matrix gives a baseline against which such advisory effects can be measured.")
    add_heading(doc, "Scientific, Technological, and Socio-Economic Impacts", 2)
    add_para(doc, "Scientifically, the project demonstrates how unsupervised wallet segmentation can be validated with network-science evidence rather than evaluated only with clustering scores. Technologically, the project creates a reusable pipeline for Ethereum wallet processing, feature extraction, clustering, statistical testing, network analysis, and temporal state modelling. Socio-economically, the work can support safer research on AI-mediated financial systems because it allows advisory, scoring, and risk mechanisms to be tested on synthetic agents derived from real behavioural patterns rather than real customers.")
    add_para(doc, "The project also has commercial and entrepreneurial relevance. Potential applications include wallet risk analysis, behavioural customer segmentation for DeFi research, market-surveillance prototyping, synthetic financial-user simulation, and explainable credit-score experimentation. However, the project is not a production credit-scoring system and should not be used to make real financial decisions without additional validation, regulatory review, and fairness analysis.")
    add_para(doc, "From a scientific perspective, the main contribution is methodological triangulation. The project does not rely on one model score. It combines clustering, feature interpretation, graph structure, non-parametric tests, transaction flow, temporal transitions, and internal subclusters. This triangulation makes the conclusion more credible. If all methods point toward compatible interpretations, confidence in the segmentation increases.")
    add_para(doc, "From a technological perspective, the project creates reusable artifacts. The scripts and CSV outputs can be rerun or extended. New months can be processed, new features can be added, and alternative clustering models can be compared. This makes the project more like a research platform than a single one-off analysis.")
    add_para(doc, "From a socio-economic perspective, the work can support better evaluation of AI in finance. Instead of testing advisory systems on real customers or purely fictional personas, future researchers can use empirically grounded synthetic agents. This reduces harm while preserving behavioural realism. It also allows fairness and explainability experiments to be repeated under controlled conditions.")
    add_heading(doc, "Freedom-to-Use and Sustainability", 2)
    add_para(doc, "No specific Freedom-to-Use issue was identified. The project relies on public blockchain-style data, open-source analytical methods, and original scripts. A full commercial deployment would still require a formal patent and license review, especially if packaged as a wallet-risk product or financial scoring tool. From a sustainability perspective, the solution is primarily analytical and offline. It does not require blockchain mining or on-chain execution. The main environmental impact is computational energy from processing large datasets. This was reduced by using monthly aggregation, sampling for graph construction, and saved intermediate CSV outputs rather than repeatedly scanning raw files.")
    add_para(doc, "The solution contributes to sustainable engineering practices by emphasizing reuse of intermediate data, modular scripts, and transparent reporting. Instead of recomputing every expensive step for every analysis, the pipeline stores clean features, cluster labels, summary tables, and figures. This reduces repeated computation and makes the work easier for another student to continue. The approach is also resource-efficient because it avoids unnecessary large-model training; the machine-learning layer uses classical unsupervised methods that are explainable and relatively lightweight.")

    doc.add_page_break()
    add_heading(doc, "ETHICAL ISSUES", 1)
    add_para(doc, "The project has several ethical dimensions even though it does not handle named personal identities or make real credit decisions. Ethereum wallet addresses are pseudonymous rather than directly personal, but behavioural analysis can still create privacy risks if combined with external information. For this reason, the report presents aggregate cluster-level findings and does not attempt to deanonymize wallet owners. The project should be understood as a research prototype, not a surveillance or enforcement system.")
    add_para(doc, "A second ethical issue is fairness. If wallet segments are later used for credit scoring or advisory decisions, there is a risk that models could penalize certain behavioural groups without sufficient explanation. This motivates the planned explainability layer and the use of transparent behavioural features. A third issue is misuse: segmentation could be used to target vulnerable users with risky financial advice. Any future advisory system should include safeguards, disclaimers, and evaluation of harmful behavioural nudging. Finally, the project avoids real financial transactions and does not issue investment advice, which limits immediate harm during development.")
    add_para(doc, "There are no identified ethical issues related to toxic materials, health hazards, or physical manufacturing. The main ethical concerns are data privacy, interpretability, fairness, and potential misuse of behavioural profiling. These concerns are acknowledged and addressed through aggregation, reproducibility, transparent reporting, and the decision to keep the work as an academic prototype.")

    add_heading(doc, "PROJECT MANAGEMENT", 1)
    add_heading(doc, "Initial and Final Plans", 2)
    add_para(doc, "The initial project plan focused on a full LLM-governed crypto-bank simulator. The intended pipeline included synthetic users, LLM advisory personas, market events, dynamic borrowing and repayment behaviour, credit scoring, and SHAP-based explainability. During implementation, the team discovered that the synthetic-user layer was the weakest assumption. Predefined personas were easy to describe but difficult to defend scientifically.")
    add_para(doc, "The final plan shifted the center of gravity to empirical behavioural segmentation. The team first processed Ethereum transaction data, engineered wallet features, performed clustering, and validated clusters statistically and graph-theoretically. This produced a stronger foundation for the later simulator. In project-management terms, the plan changed from building the whole simulator breadth-first to building the data and modelling foundation depth-first.")
    add_heading(doc, "What Changed and Why", 2)
    add_bullets(doc, [
        "The project moved from persona-driven agent definitions to segmentation-driven agent definitions because empirical wallet behaviour is more defensible than manually assigned personas.",
        "The implementation expanded network-science validation because clustering metrics alone were not enough to prove that the groups were meaningful.",
        "The temporal state model was added to capture month-to-month behavioural changes, which the original static persona design could not represent.",
        "The LLM-advisory and SHAP components were moved to future work after the segmentation layer became the core scientific contribution.",
    ])
    add_heading(doc, "Lessons Learned", 2)
    add_para(doc, "The main project-management lesson is that realistic data modelling can change the correct implementation order. The original vision was valid, but its first step needed to be revised. Building a simulator before validating user types would have produced a system that looked complete but rested on weak assumptions. The team also learned that large datasets require modular scripts, intermediate artifacts, and careful separation between raw processing, feature engineering, modelling, validation, and reporting. Finally, the project showed the value of producing interpretable evidence at every stage: cluster labels, statistical tests, graphs, figures, and transition matrices all helped convert raw outputs into an explainable engineering result.")

    add_heading(doc, "CONCLUSION AND FUTURE WORK", 1)
    add_para(doc, "The project successfully produced a data-driven Ethereum wallet behaviour segmentation and validated it through multiple independent forms of evidence. KMeans6 clusters are interpretable as behavioural roles, statistically separable across core features, strongly aligned with feature-similarity graph communities, meaningful in raw transaction networks, and useful as temporal behavioural states. The most important quantitative findings are the 0.943 intra-cluster edge ratio, 0.777 modularity, 0.932 assortativity, 5.66 same-cluster lift over random mixing, and permutation-test p-value of 0.000999. These results support the conclusion that the clusters are not arbitrary algorithmic labels but meaningful wallet behaviour groups.")
    add_para(doc, "The main limitation is that the final project emphasizes segmentation and validation more than complete end-to-end simulation. The LLM advisory, dynamic credit-score, and SHAP explainability components remain logical next steps. Another limitation is that raw transaction coverage was partial for transaction-network analysis, although coverage around 75% is strong enough for cluster-level interpretation. The analysis is also based on engineered features and selected temporal windows; future work should test robustness across additional datasets, more months, token categories, and market regimes.")
    add_para(doc, "Future work should integrate the validated clusters into the original simulator architecture. Each cluster can initialize a class of agents with realistic activity, value-flow, volatility, and transition tendencies. Market events can then perturb monthly states, advisory personas can influence agent decisions, and SHAP or similar methods can explain credit-score changes. Additional improvements include graph embeddings, temporal graph neural networks, persistent homology for subcluster validation, dashboard visualization, fairness audits, and energy-efficient incremental processing. For real-world implementation, the system would need stronger data governance, license review, model-risk controls, and safeguards against harmful behavioural targeting.")

    doc.add_page_break()
    add_heading(doc, "APPENDIX", 1)
    add_heading(doc, "Appendix A: Main Output Files", 2)
    add_table(doc, ["Artifact", "Purpose"], [
        ("dataverse_files/wallet_kmeans_clusters.csv", "Authoritative KMeans6 wallet cluster assignments."),
        ("dataverse_files/wallet_kmeans6_expanded_daily_features_robust_scaled.csv", "Expanded robust-scaled behavioural feature set."),
        ("dataverse_files/network_science_kmeans6_full/broad_network_science_summary.md", "Summary of cluster interpretation, graph validation, and transaction-network evidence."),
        ("dataverse_files/network_science_kmeans6_full/mathematical_proofs_and_statistics.md", "Formal notation, proofs, and statistical evidence for cluster validity."),
        ("dataverse_files/cluster_transition_analysis_expanded/transitional_scientific_report.md", "Monthly behavioural state transition analysis."),
        ("dataverse_files/cluster_behavior_subclusters_expanded/behavioral_subcluster_scientific_report.md", "Internal subcluster analysis within the six main clusters."),
    ], widths=[3.5, 3.0])
    add_heading(doc, "Appendix B: Key Scripts", 2)
    add_table(doc, ["Script", "Role"], [
        ("build_full_kmeans6_expanded_features.py", "Builds expanded wallet feature dataset."),
        ("analysis_kmean6.py", "Runs and inspects KMeans6 clustering results."),
        ("broad_network_science_analysis.py", "Creates graph-theoretic and transaction-network analyses."),
        ("feature_graph_randomness_test.py", "Runs random-label permutation test for feature graph alignment."),
        ("cluster_transition_analysis.py", "Builds monthly state transition matrices."),
        ("cluster_behavior_subcluster_analysis.py", "Analyses internal behavioural subclusters."),
    ], widths=[3.1, 3.4])

    add_heading(doc, "REFERENCES", 1)
    refs = [
        "Byrd, D., Hybinette, M., & Balch, T. H. (2019). ABIDES: Towards high-fidelity market simulation for AI research. arXiv:1904.12066.",
        "Chen, B., Wu, Z., & Zhao, R. (2023). From fiction to fact: The growing role of generative AI in business and finance. Journal of Chinese Economic and Business Studies, 21(4), 471-496.",
        "Dao, M. (2022). The Maker Protocol: MakerDAO's multi-collateral Dai system.",
        "Diederich, S., Brendel, A. B., & Kolbe, L. M. (2020). Designing anthropomorphic enterprise conversational agents. Business & Information Systems Engineering, 62(3), 193-209.",
        "European Union. (2016). General Data Protection Regulation (GDPR), Regulation (EU) 2016/679. Official Journal of the European Union.",
        "ISO/IEC. (2011). ISO/IEC 25010:2011, Systems and software engineering - Systems and software Quality Requirements and Evaluation.",
        "Mazzarisi, P., Ravagnani, A., Deriu, P., Lillo, F., Medda, F., & Russo, A. (2022). A machine learning approach to support decision in insider trading detection. arXiv:2212.05912.",
        "Newman, M. E. J. (2010). Networks: An Introduction. Oxford University Press.",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "Schmitt, M. (2024). Explainable automated machine learning for credit decisions: Enhancing human-artificial intelligence collaboration in financial engineering. arXiv:2402.03806.",
        "Zaharia, M., Chen, A., Davidson, A., Ghodsi, A., Hong, S., Konwinski, A., et al. (2018). MLflow: A platform for ML development and reproducibility. Proceedings of SysML.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.add_run(ref).font.name = "Arial"
        p.runs[0].font.size = Pt(10)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
