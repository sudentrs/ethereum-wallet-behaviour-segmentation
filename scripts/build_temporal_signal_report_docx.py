from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path("results/full_eoa_temporal_signal_streaming_k5/model_selection")
OUT = BASE / "temporal_signal_results_report_en.docx"
FOLLOWUP = BASE / "followup_analyses"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 31, 31)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "E8EEF5"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tbl_cell_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("Temporal Signal Results Report")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Time-domain and wavelength evidence for the EOA-only Behavior K5 taxonomy")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, [6.5])
    set_cell_margins(table, top=120, bottom=120, start=160, end=160)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    r.font.size = Pt(11)
    p = cell.add_paragraph(body)
    p.paragraph_format.space_after = Pt(0)
    for r in p.runs:
        r.font.size = Pt(10.5)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    set_cell_margins(table)
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_shading(hdr[i], LIGHT_FILL)
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = INK
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 4 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            run.font.color.rgb = INK
    doc.add_paragraph()
    return table


def pct(x):
    return f"{float(x) * 100:.2f}%"


def num(x, n=4):
    return f"{float(x):.{n}f}"


def load_data():
    time_ms = pd.read_csv(BASE / "time_domain_model_selection_summary.csv")
    wave_ms = pd.read_csv(BASE / "wavelength_model_selection_summary.csv")
    time_profiles = pd.read_csv(BASE / "time_domain_k5_cluster_profiles.csv")
    time_labels = pd.read_csv(BASE / "time_domain_k5_interpretive_labels.csv")
    wave_profiles = pd.read_csv(BASE / "wavelength_k3_cluster_profiles.csv")
    time_rel = pd.read_csv(BASE / "behavior_k5_vs_time_domain_cluster_long_summary.csv")
    wave_rel = pd.read_csv(BASE / "behavior_k5_vs_wavelength_cluster_long_summary.csv")
    return time_ms, wave_ms, time_profiles, time_labels, wave_profiles, time_rel, wave_rel


def load_followup_data():
    transition_summary = pd.read_csv(FOLLOWUP / "temporal_transition_predictiveness_summary.csv")
    transition_by_state = pd.read_csv(FOLLOWUP / "temporal_transition_predictiveness_by_state.csv")
    time_outcomes = pd.read_csv(FOLLOWUP / "behavior_k5_time_domain_outcome_profile.csv")
    trajectory_bridge = pd.read_csv(FOLLOWUP / "trajectory_type_signal_bridge_by_behavior_k5.csv")
    return transition_summary, transition_by_state, time_outcomes, trajectory_bridge


def main():
    time_ms, wave_ms, time_profiles, time_labels, wave_profiles, time_rel, wave_rel = load_data()

    doc = Document()
    style_document(doc)
    add_title(doc)

    add_callout(
        doc,
        "Analytical summary",
        "The EOA-only Behavior K5 taxonomy is retained as the main wallet-role structure. The temporal signal layer is represented with time-domain K5 for activity-state interpretation and wavelength K3 for rhythm interpretation. Time-domain K3 provides a compact alternative, but it collapses most wallets into one dominant sparse state and shows weaker seed stability.",
    )

    doc.add_heading("1. Time-Domain Signal Layer", level=1)
    doc.add_paragraph(
        "The time-domain layer measures activity continuity, inactivity gaps, burstiness, active/inactive switching, and short-term transaction changes. It is therefore the right validation layer for temporal persistence and activity-state behavior rather than transaction-network homophily."
    )

    td_rows = []
    for _, r in time_ms.iterrows():
        td_rows.append(
            [
                int(r["k"]),
                num(r["silhouette_sample"]),
                num(r["seed_stability_mean_ari"]),
                pct(r["cluster_share_min"]),
                pct(r["cluster_share_max"]),
                num(r["cluster_share_entropy_norm"]),
                num(r["selection_score"]),
            ]
        )
    add_table(
        doc,
        ["K", "Silhouette", "Mean ARI", "Min share", "Max share", "Entropy", "Score"],
        td_rows,
        [0.45, 0.95, 0.85, 0.85, 0.85, 0.75, 0.8],
    )

    doc.add_heading("Time-domain decision", level=2)
    doc.add_paragraph(
        "The automatic ranking selects K3 because it has the strongest silhouette score. However, K3 places 76.41% of wallets into one cluster and has weak seed stability (mean ARI = 0.2618). This structure is compact, but it is too coarse for an explanatory temporal-state taxonomy."
    )
    doc.add_paragraph(
        "K5 is more defensible as the main time-domain state representation. It has materially stronger stability (mean ARI = 0.5591), a much more balanced largest cluster (36.57%), and the highest balance/entropy score among the tested time-domain options. It also creates interpretable states that can be connected back to Behavior K5."
    )

    doc.add_heading("2. Time-Domain K5 Profiles", level=1)
    label_map = dict(zip(time_labels["cluster"], time_labels["interpretive_label"]))
    profile_rows = []
    for _, r in time_profiles.iterrows():
        c = int(r["cluster"])
        profile_rows.append(
            [
                f"TD{c}",
                label_map.get(c, ""),
                pct(r["wallet_share"]),
                f"{float(r['active_days']):.2f}",
                f"{float(r['longest_inactive_gap']):.2f}",
                f"{float(r['active_inactive_switches']):.2f}",
                f"{float(r['tx_burstiness_cv']):.3f}",
            ]
        )
    add_table(
        doc,
        ["Cluster", "Interpretive label", "Wallet share", "Active days", "Inactive gap", "Switches", "Burstiness"],
        profile_rows,
        [0.6, 1.75, 0.85, 0.75, 0.85, 0.75, 0.8],
        font_size=8,
    )
    add_callout(
        doc,
        "Most important time-domain state",
        "TD4 is the clearest intermittent switching state. It contains only 3.81% of wallets, but it has far more active days, many more active/inactive switches, higher burstiness, and a much higher max-to-mean transaction ratio than the sparse long-gap states.",
    )

    doc.add_heading("3. Wavelength Signal Layer", level=1)
    doc.add_paragraph(
        "The wavelength layer captures spectral rhythm: dominant period, low/medium/high-frequency power shares, weekly and monthly bands, peak concentration, and spectral entropy. This layer describes rhythm and temporal context structure rather than causal event effects."
    )
    wave_rows = []
    for _, r in wave_ms.iterrows():
        wave_rows.append(
            [
                int(r["k"]),
                num(r["silhouette_sample"]),
                num(r["seed_stability_mean_ari"]),
                pct(r["cluster_share_min"]),
                pct(r["cluster_share_max"]),
                num(r["cluster_share_entropy_norm"]),
                num(r["selection_score"]),
            ]
        )
    add_table(
        doc,
        ["K", "Silhouette", "Mean ARI", "Min share", "Max share", "Entropy", "Score"],
        wave_rows,
        [0.45, 0.95, 0.85, 0.85, 0.85, 0.75, 0.8],
    )

    doc.add_heading("Wavelength decision", level=2)
    doc.add_paragraph(
        "Wavelength K3 is the cleanest signal-model selection result. It has the highest selection score, the best silhouette among the tested wavelength solutions, and very strong seed stability (mean ARI = 0.8317). Its cluster sizes are also acceptable, with the smallest rhythm group still covering 6.41% of wallets."
    )

    doc.add_heading("4. Wavelength K3 Profiles", level=1)
    wave_profile_rows = []
    for _, r in wave_profiles.iterrows():
        c = int(r["cluster"])
        label = (
            "short/medium distributed high-frequency rhythm"
            if c == 0
            else "longer distributed rhythm"
            if c == 1
            else "slow-cycle concentrated rhythm"
        )
        wave_profile_rows.append(
            [
                f"W{c}",
                label,
                pct(r["wallet_share"]),
                f"{float(r['tx_dominant_period_days']):.2f}",
                f"{float(r['tx_high_frequency_power_share']):.3f}",
                f"{float(r['tx_low_frequency_power_share']):.3f}",
                f"{float(r['tx_spectral_entropy']):.3f}",
            ]
        )
    add_table(
        doc,
        ["Cluster", "Rhythm interpretation", "Wallet share", "Dominant period", "TX high", "TX low", "Entropy"],
        wave_profile_rows,
        [0.6, 2.0, 0.8, 0.95, 0.65, 0.65, 0.65],
        font_size=8,
    )
    add_callout(
        doc,
        "Most important wavelength state",
        "W2 is the most distinctive rhythm group. It has a long dominant period of about 342.5 days, a much higher low-frequency power share, and lower spectral entropy, which indicates a more concentrated slow-cycle rhythm.",
    )

    doc.add_heading("5. Relationship with Behavior K5", level=1)
    doc.add_paragraph(
        "The signal clusters do not simply reproduce the Behavior K5 labels, but they are also not uniformly distributed across them. The main K5 behavioral roles therefore carry different temporal signatures."
    )

    top_time = (
        time_rel.sort_values(["behavior_k5_cluster", "within_behavior_share"], ascending=[True, False])
        .groupby("behavior_k5_cluster")
        .head(2)
    )
    time_relationship_rows = []
    for _, r in top_time.iterrows():
        time_relationship_rows.append(
            [
                int(r["behavior_k5_cluster"]),
                f"TD{int(r['time_domain_cluster'])}",
                r["label"],
                pct(r["within_behavior_share"]),
                f"{float(r['enrichment_vs_global']):.2f}x",
            ]
        )
    add_table(
        doc,
        ["Behavior K5", "TD state", "Label", "Within share", "Enrichment"],
        time_relationship_rows,
        [0.85, 0.65, 2.45, 1.0, 0.9],
        font_size=8,
    )

    top_wave = (
        wave_rel.sort_values(["behavior_k5_cluster", "within_behavior_share"], ascending=[True, False])
        .groupby("behavior_k5_cluster")
        .head(2)
    )
    wave_relationship_rows = []
    for _, r in top_wave.iterrows():
        c = int(r["wavelength_cluster"])
        label = "slow-cycle concentrated rhythm" if c == 2 else "distributed rhythm"
        wave_relationship_rows.append(
            [
                int(r["behavior_k5_cluster"]),
                f"W{c}",
                label,
                pct(r["within_behavior_share"]),
                f"{float(r['enrichment_vs_global']):.2f}x",
            ]
        )
    add_table(
        doc,
        ["Behavior K5", "W rhythm", "Rhythm interpretation", "Within share", "Enrichment"],
        wave_relationship_rows,
        [0.85, 0.65, 2.45, 1.0, 0.9],
        font_size=8,
    )

    doc.add_heading("6. Main Findings", level=1)
    add_bullets(
        doc,
        [
            "Behavior K5 cluster 3 is strongly enriched in the intermittent switching time-domain state (5.32x over the global expectation).",
            "Behavior K5 cluster 3 is also strongly enriched in the slow-cycle/concentrated wavelength state (4.55x).",
            "Behavior K5 cluster 2 is strongly concentrated in the sparsest long-gap time-domain state (3.49x).",
            "Behavior K5 cluster 4 is mostly low-activity mixed and sparse long-gap, but still shows a meaningful slow-cycle rhythm signal (1.85x enrichment in W2).",
            "The temporal signal layer supports Behavior K5 as a behavioral taxonomy with time-structured signatures, not merely a static transaction-intensity partition.",
        ],
    )

    doc.add_heading("7. Additional Validation Analyses", level=1)
    transition_summary, transition_by_state, time_outcomes, trajectory_bridge = load_followup_data()
    ts = transition_summary.iloc[0]
    add_callout(
        doc,
        "Temporal transition predictiveness",
        f"Across {int(ts['total_wallet_month_pairs']):,} consecutive wallet-month pairs, the overall self-transition rate is {100 * ts['overall_self_transition_rate']:.1f}%. Conditioning on the current monthly state reduces next-state entropy by {100 * ts['normalized_entropy_reduction']:.1f}%, indicating moderate next-month predictiveness.",
    )

    transition_rows = []
    for _, r in transition_by_state.iterrows():
        transition_rows.append(
            [
                int(r["from_state"]),
                f"{100 * r['self_transition_probability']:.1f}%",
                int(r["top_next_state"]),
                f"{100 * r['top_next_probability']:.1f}%",
                f"{r['top_next_lift_vs_marginal']:.2f}x",
            ]
        )
    add_table(
        doc,
        ["Current state", "Self-transition", "Top next", "Top-next prob.", "Lift vs marginal"],
        transition_rows,
        [1.0, 1.15, 0.9, 1.2, 1.2],
        font_size=8.5,
    )

    doc.add_heading("Behavior K5 x Signal Outcome Profiling", level=2)
    doc.add_paragraph(
        "The outcome profiling compares temporal activity outcomes within each Behavior K5 role after signal clusters are assigned. The table below reports the dominant time-domain state within each Behavior K5 role."
    )
    dominant_td = (
        time_outcomes.sort_values(["k_cluster", "within_behavior_share"], ascending=[True, False])
        .groupby("k_cluster")
        .head(1)
    )
    dominant_td_rows = []
    for _, r in dominant_td.iterrows():
        dominant_td_rows.append(
            [
                int(r["k_cluster"]),
                f"TD{int(r['time_domain_cluster'])}",
                f"{100 * r['within_behavior_share']:.1f}%",
                f"{r['active_days_mean']:.2f}",
                f"{r['active_inactive_switches_mean']:.2f}",
                f"{r['longest_inactive_gap_mean']:.2f}",
                f"{r['tx_burstiness_cv_mean']:.3f}",
            ]
        )
    add_table(
        doc,
        ["Behavior K5", "Dominant TD", "Share", "Active days", "Switches", "Inactive gap", "Burstiness"],
        dominant_td_rows,
        [0.85, 0.8, 0.75, 0.85, 0.75, 0.9, 0.85],
        font_size=8,
    )

    doc.add_paragraph(
        "The outcome profile is consistent with the crosstab evidence. Behavior K5 cluster 2 is dominated by the sparsest long-gap state, with approximately 3.30 active days and a mean inactive gap of 289.71 days in its dominant TD1 segment. Behavior K5 cluster 4 is dominated by TD3 and shows a more active mixed pattern, with approximately 32.80 active days and 47.49 active/inactive switches."
    )

    doc.add_heading("Returner and One-Time Switcher Bridge", level=2)
    doc.add_paragraph(
        "The trajectory bridge links the existing monthly K5 trajectory taxonomy to the signal evidence. The bridge is computed by starting Behavior K5 state, because the retained repository outputs contain trajectory summaries by start state rather than wallet-level trajectory assignments."
    )
    bridge_rows = []
    for _, r in trajectory_bridge.iterrows():
        bridge_rows.append(
            [
                int(r["behavior_k5_start_state"]),
                f"{100 * r['returner']:.1f}%",
                f"{100 * r['one_time_switcher']:.1f}%",
                f"{r['td4_intermit_switching_enrichment']:.2f}x",
                f"{r['w2_slow_cycle_enrichment']:.2f}x",
            ]
        )
    add_table(
        doc,
        ["Behavior K5 start", "Returner", "One-time switcher", "TD4 enrichment", "W2 enrichment"],
        bridge_rows,
        [1.05, 0.9, 1.1, 1.15, 1.0],
        font_size=8.5,
    )
    add_callout(
        doc,
        "Trajectory evidence",
        "Behavior K5 cluster 3 combines a high one-time-switcher share with strong intermittent-switching enrichment (5.32x) and slow-cycle wavelength enrichment (4.55x). Behavior K5 cluster 4 is more returner-oriented (51.5%) but still carries intermittent and slow-cycle signal enrichment.",
    )

    doc.add_heading("8. Synthesis", level=1)
    doc.add_paragraph(
        "The temporal interpretation of the EOA-only K5 wallet taxonomy is supported by two complementary temporal signal spaces: time-domain activity features and wavelength/spectral rhythm features. The time-domain analysis shows that a richer K5 state representation captures sparse long-gap, low-activity mixed, and intermittent switching regimes. The wavelength analysis supports a compact K3 rhythm structure, separating wallets into distributed high-frequency and concentrated slow-cycle rhythm profiles."
    )
    doc.add_paragraph(
        "Cross-tabulation with the main Behavior K5 taxonomy shows that these temporal signals are not uniformly distributed across wallet roles. Behavior K5 cluster 3 is strongly enriched in intermittent switching states and slow-cycle wavelength patterns, while Behavior K5 cluster 2 is strongly enriched in sparse long-gap states. These results support the interpretation that the main wallet roles encode not only static transaction intensity but also distinct temporal activity regimes."
    )
    add_callout(
        doc,
        "Main conclusion",
        "Behavior K5 remains the main wallet-role taxonomy. Time-domain K5 and wavelength K3 provide complementary temporal evidence that the K5 roles are associated with distinct activity-state and rhythm profiles. The transition and trajectory analyses show moderate next-month predictiveness and a link to returner/one-time-switcher behavior. The evidence is behavioral and associational rather than causal.",
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Temporal signal analysis | EOA-only Behavior K5")

    doc.core_properties.title = "Temporal Signal Results Report"
    doc.core_properties.subject = "Time-domain and wavelength evidence for EOA-only Behavior K5"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
